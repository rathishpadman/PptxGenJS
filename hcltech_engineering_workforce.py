"""HCLTech slide: Engineering Workforce — 1,000+ and compounding."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TEMPLATE = '/root/.claude/skills/synced/hcltech-deck/assets/hcltech_base.pptx'
OUT = '/home/user/PptxGenJS/hcltech_engineering_workforce.pptx'

BLACK       = RGBColor(0x00, 0x00, 0x00)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
PURPLE      = RGBColor(0x5F, 0x1E, 0xBE)
PURPLE_DEEP = RGBColor(0x41, 0x14, 0x82)
BLUE        = RGBColor(0x0F, 0x5F, 0xDC)
BRIGHT_BLUE = RGBColor(0x3C, 0x91, 0xFF)
ICE_BLUE    = RGBColor(0xDC, 0xE6, 0xF0)
GREY_1      = RGBColor(0x82, 0x91, 0xA0)
GREY_3      = RGBColor(0xC8, 0xD2, 0xDD)
GREY_4      = RGBColor(0xE6, 0xEB, 0xF5)
FONT = 'Aptos'

prs = Presentation(TEMPLATE)
slide = prs.slides.add_slide(prs.slide_layouts[16])
S = slide.shapes
I = Inches


def style(run, size, *, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def fill_tf(tf, paras, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            line_spacing=None, space_after=0):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, prun in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for text, size, bold, color in prun:
            style(p.add_run(), size, bold=bold, color=color)
            p.runs[-1].text = text


def textbox(x, y, w, h, runs, **kw):
    tb = S.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, runs if runs and isinstance(runs[0], list) else [runs], **kw)
    return tb


def shape(kind, x, y, w, h, *, fill=None, line=None, line_w=0.75, adj=None):
    s = S.add_shape(kind, I(x), I(y), I(w), I(h))
    if adj is not None:
        vals = adj if isinstance(adj, (list, tuple)) else [adj]
        for i, v in enumerate(vals):
            if i < len(s.adjustments):
                s.adjustments[i] = v
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    return s


def poly(pts, color, width, dash=False):
    b = S.build_freeform(I(pts[0][0]), I(pts[0][1]))
    b.add_line_segments([(I(x), I(y)) for x, y in pts[1:]], close=False)
    s = b.convert_to_shape()
    s.fill.background()
    s.line.color.rgb = color
    s.line.width = Pt(width)
    if dash:
        s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    s.shadow.inherit = False
    return s


def arrow(cx, cy, color=GREY_1, size=0.09):
    poly([(cx - size, cy), (cx + size, cy)], color, 1.1)
    poly([(cx + size * .45, cy - size * .6), (cx + size, cy),
          (cx + size * .45, cy + size * .6)], color, 1.1)


def glyph(name, cx, cy, s, color):
    u = s / 2.0
    lw = max(1.0, s * 2.2)
    if name == 'gear':
        shape(MSO_SHAPE.GEAR_6, cx - u * .86, cy - u * .86, s * .86, s * .86,
              line=color, line_w=lw)
        shape(MSO_SHAPE.OVAL, cx - u * .24, cy - u * .24, s * .24, s * .24,
              line=color, line_w=lw)
    elif name == 'nodes':
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - s * .10, cy - u * .74,
              s * .20, s * .20, fill=color, adj=0.25)
        for dx in (-.32, -.08, .16):
            shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + s * dx, cy + u * .28,
                  s * .16, s * .16, fill=color, adj=0.25)
        poly([(cx - s * .24, cy + u * .08), (cx + s * .24, cy + u * .08)],
             color, lw * .85)
    elif name == 'cubes':
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - s * .15, cy - u * .82,
              s * .30, s * .30, fill=color, adj=0.22)
        for dx in (-.38, .08):
            shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + s * dx, cy + u * .04,
                  s * .30, s * .30, fill=color, adj=0.22)
    elif name == 'database':
        for i in range(3):
            shape(MSO_SHAPE.OVAL, cx - u * .70, cy - u * .82 + i * s * .32,
                  s * .70, s * .30, line=color, line_w=lw)
    elif name == 'people':
        for dx in (-u * .44, u * .10):
            shape(MSO_SHAPE.OVAL, dx + cx, cy - u * .74, s * .30, s * .30,
                  line=color, line_w=lw)
            shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, dx + cx - s * .05,
                  cy - u * .12, s * .40, s * .44, line=color, line_w=lw,
                  adj=(0.45, 0.0))
    elif name == 'arrow_up':
        shape(MSO_SHAPE.UP_ARROW, cx - u * .8, cy - u * .8, s * .8, s * .8,
              fill=color, adj=(0.42, 0.42))
        S.shapes[-1] if False else None
    elif name == 'cycle':
        shape(MSO_SHAPE.OVAL, cx - u * .74, cy - u * .74, s * .74, s * .74,
              line=color, line_w=lw)
        shape(MSO_SHAPE.ISOSCELES_TRIANGLE, cx + u * .18, cy - u * .60,
              s * .30, s * .26, fill=color)
    elif name == 'chart_up':
        for i, hh in enumerate((.34, .58, .84)):
            shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - u * .70 + i * s * .27,
                  cy + u * .62 - s * hh, s * .18, s * hh, fill=color, adj=0.3)


# ------------------------------------------------------------------ header ---
textbox(0.65, 0.50, 12.03, 0.55,
        [[('Engineering workforce: ', 28, True, BLACK),
          ('1,000+ and compounding', 28, True, PURPLE)]])
textbox(0.65, 1.12, 12.03, 0.32,
        [('An AI-native engineering engine — capability compounds because the '
          'IP compounds.', 16, False, GREY_1)])

# -------------------------------------------------------------- hero figure ---
HW, HH, HY = 4.10, 0.84, 1.62
HX = (13.333 - HW) / 2
hero = shape(MSO_SHAPE.ROUNDED_RECTANGLE, HX, HY, HW, HH, adj=0.14)
hero.fill.gradient()
hero.fill.gradient_stops[0].color.rgb = PURPLE_DEEP
hero.fill.gradient_stops[0].position = 0.0
hero.fill.gradient_stops[1].color.rgb = PURPLE
hero.fill.gradient_stops[1].position = 1.0
hero.fill.gradient_angle = 0.0
shape(MSO_SHAPE.OVAL, HX + 0.36, HY + 0.16, 0.52, 0.52, fill=WHITE)
glyph('people', HX + 0.62, HY + 0.42, 0.30, PURPLE)
textbox(HX + 1.06, HY, HW - 1.30, HH,
        [[('1,000+', 30, True, WHITE)],
         [('Engineering base', 12, False, ICE_BLUE)]],
        anchor=MSO_ANCHOR.MIDDLE, space_after=1)

# ------------------------------------------------------------ pillar cards ---
CW, CGAP, CT, CH = 2.82, 0.25, 3.14, 2.20
PX = [0.65 + i * (CW + CGAP) for i in range(4)]
PILLARS = [
    ('gear',     'Agentic &\nAutomation',    '400+', PURPLE_DEEP),
    ('nodes',    'AI &\nGen AI',             '200+', PURPLE),
    ('cubes',    'Product &\nPlatform',      '200+', BLUE),
    ('database', 'Data &\nProcess Mining',   '200+', BRIGHT_BLUE),
]

# connector tree from the hero down to each card
BUS_Y = 2.82
poly([(13.333 / 2, HY + HH), (13.333 / 2, BUS_Y)], GREY_3, 1.0)
poly([(PX[0] + CW / 2, BUS_Y), (PX[3] + CW / 2, BUS_Y)], GREY_3, 1.0)
for x in PX:
    poly([(x + CW / 2, BUS_Y), (x + CW / 2, CT)], GREY_3, 1.0)
    shape(MSO_SHAPE.OVAL, x + CW / 2 - 0.05, BUS_Y - 0.05, 0.10, 0.10, fill=GREY_3)

for (icon, name, count, col), x in zip(PILLARS, PX):
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, CT, CW, CH,
          fill=WHITE, line=GREY_3, adj=0.035)
    shape(MSO_SHAPE.OVAL, x + CW / 2 - 0.29, CT + 0.18, 0.58, 0.58, fill=col)
    glyph(icon, x + CW / 2, CT + 0.47, 0.30, WHITE)
    textbox(x + 0.18, CT + 0.88, CW - 0.36, 0.46,
            [[(ln, 12.5, True, PURPLE_DEEP)] for ln in name.split('\n')],
            align=PP_ALIGN.CENTER, line_spacing=0.95)
    poly([(x + 0.55, CT + 1.42), (x + CW - 0.55, CT + 1.42)], GREY_4, 0.75)
    textbox(x, CT + 1.50, CW, 0.48, [(count, 28, True, col)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(x, CT + 2.00, CW, 0.20, [('People', 10, False, GREY_1)],
            align=PP_ALIGN.CENTER)

# ---------------------------------------------------------- the flywheel -----
BY, BH = 5.52, 1.44
shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, BY, 12.03, BH, fill=GREY_4, adj=0.09)
textbox(0.65, BY + 0.10, 12.03, 0.22,
        [('THE COMPOUNDING FLYWHEEL', 9.5, True, PURPLE)], align=PP_ALIGN.CENTER)

STEPS = [('people', '1. Elite Talent'), ('arrow_up', '2. Client Deployments'),
         ('cubes', '3. Reusable Agent IP'), ('cycle', '4. Faster Delivery'),
         ('chart_up', '5. More Client Wins')]
INNER_X, INNER_W = 1.00, 11.33
step_w = INNER_W / len(STEPS)
CIRC_Y = BY + 0.62
for i, (icon, name) in enumerate(STEPS):
    cx = INNER_X + step_w / 2 + i * step_w
    shape(MSO_SHAPE.OVAL, cx - 0.24, CIRC_Y - 0.24, 0.48, 0.48,
          fill=WHITE, line=GREY_3)
    glyph(icon, cx, CIRC_Y, 0.26, PURPLE)
    textbox(cx - step_w / 2, CIRC_Y + 0.32, step_w, 0.22,
            [(name, 9.5, False, BLACK)], align=PP_ALIGN.CENTER)
    if i < len(STEPS) - 1:
        arrow(cx + step_w / 2, CIRC_Y, PURPLE)

# loop back: wins feed talent
LX_, RX_ = INNER_X + step_w / 2, INNER_X + step_w / 2 + 4 * step_w
LOOP_Y = BY + BH - 0.16
poly([(RX_, CIRC_Y + 0.60), (RX_, LOOP_Y), (LX_, LOOP_Y), (LX_, CIRC_Y + 0.60)],
     PURPLE, 1.0, dash=True)
poly([(LX_ - 0.055, CIRC_Y + 0.66), (LX_, CIRC_Y + 0.60),
      (LX_ + 0.055, CIRC_Y + 0.66)], PURPLE, 1.1)

prs.save(OUT)
print('saved', OUT, '| shapes:', len(slide.shapes))
