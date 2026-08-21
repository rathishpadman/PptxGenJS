"""HCLTech slide: Re-energizing our GTM — three plays, one motion."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TEMPLATE = '/root/.claude/skills/synced/hcltech-deck/assets/hcltech_base.pptx'
OUT = '/home/user/PptxGenJS/hcltech_gtm_three_plays.pptx'

BLACK       = RGBColor(0x00, 0x00, 0x00)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
PURPLE      = RGBColor(0x5F, 0x1E, 0xBE)
PURPLE_DEEP = RGBColor(0x41, 0x14, 0x82)
BLUE        = RGBColor(0x0F, 0x5F, 0xDC)
ICE_BLUE    = RGBColor(0xDC, 0xE6, 0xF0)
GREY_1      = RGBColor(0x82, 0x91, 0xA0)
GREY_3      = RGBColor(0xC8, 0xD2, 0xDD)
GREY_4      = RGBColor(0xE6, 0xEB, 0xF5)
FONT = 'Aptos'

prs = Presentation(TEMPLATE)
slide = prs.slides.add_slide(prs.slide_layouts[16])
S = slide.shapes
I = Inches


def style(run, size, *, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def fill_tf(tf, paras, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            line_spacing=None, space_after=0):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, prun in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for text, size, bold, color in prun:
            style(p.add_run(), size, bold=bold, color=color)
            p.runs[-1].text = text


def textbox(x, y, w, h, runs, **kw):
    tb = S.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, runs if runs and isinstance(runs[0], list) else [runs], **kw)
    return tb


def shape(kind, x, y, w, h, *, fill=None, line=None, line_w=0.75, adj=None, rot=None):
    s = S.add_shape(kind, I(x), I(y), I(w), I(h))
    if adj is not None:
        vals = adj if isinstance(adj, (list, tuple)) else [adj]
        for i, v in enumerate(vals):
            if i < len(s.adjustments):
                s.adjustments[i] = v
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    if rot is not None:
        s.rotation = rot
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    return s


def poly(pts, color, width):
    b = S.build_freeform(I(pts[0][0]), I(pts[0][1]))
    b.add_line_segments([(I(x), I(y)) for x, y in pts[1:]], close=False)
    s = b.convert_to_shape()
    s.fill.background()
    s.line.color.rgb = color
    s.line.width = Pt(width)
    s.shadow.inherit = False
    return s


def chip(kind, x, y, w, h, lines, *, fill, line=None, adj=0.5, align=PP_ALIGN.CENTER,
         pad=0.08):
    c = shape(kind, x, y, w, h, fill=fill, line=line, adj=adj)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = I(pad)
    tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, lines, align=align, anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.98)
    return c


def glyph(name, cx, cy, s, color):
    u = s / 2.0
    lw = max(1.0, s * 2.2)
    if name == 'target':
        shape(MSO_SHAPE.OVAL, cx - u, cy - u, s, s, line=color, line_w=lw)
        shape(MSO_SHAPE.OVAL, cx - u * .55, cy - u * .55, s * .55, s * .55,
              line=color, line_w=lw)
        shape(MSO_SHAPE.OVAL, cx - u * .18, cy - u * .18, s * .18, s * .18, fill=color)
    elif name == 'people':
        for dx in (-u * .44, u * .10):
            shape(MSO_SHAPE.OVAL, dx + cx, cy - u * .74, s * .30, s * .30,
                  line=color, line_w=lw)
            shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, dx + cx - s * .05,
                  cy - u * .12, s * .40, s * .44, line=color, line_w=lw,
                  adj=(0.45, 0.0))
    elif name == 'chart_up':
        for i, hh in enumerate((.34, .58, .84)):
            shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - u * .70 + i * s * .27,
                  cy + u * .62 - s * hh, s * .18, s * hh, fill=color, adj=0.3)
    elif name == 'cloud':
        shape(MSO_SHAPE.CLOUD, cx - u * .90, cy - u * .66, s * .90, s * .66,
              line=color, line_w=lw)
    elif name == 'arrow_up':
        shape(MSO_SHAPE.UP_ARROW, cx - u * .74, cy - u * .74, s * .74, s * .74,
              fill=color, adj=(0.42, 0.42))
    elif name == 'nodes':
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - s * .10, cy - u * .74,
              s * .20, s * .20, fill=color, adj=0.25)
        for dx in (-.32, -.08, .16):
            shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + s * dx, cy + u * .28,
                  s * .16, s * .16, fill=color, adj=0.25)
        poly([(cx - s * .24, cy + u * .08), (cx + s * .24, cy + u * .08)],
             color, lw * .85)
    elif name == 'database':
        for i in range(3):
            shape(MSO_SHAPE.OVAL, cx - u * .70, cy - u * .82 + i * s * .32,
                  s * .70, s * .30, line=color, line_w=lw)
    elif name == 'cubes':
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - s * .15, cy - u * .82,
              s * .30, s * .30, fill=color, adj=0.22)
        for dx in (-.38, .08):
            shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + s * dx, cy + u * .04,
                  s * .30, s * .30, fill=color, adj=0.22)
    elif name == 'doc':
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - u * .62, cy - u * .86,
              s * .62, s * .86, line=color, line_w=lw, adj=0.12)
        for i in range(3):
            poly([(cx - u * .34, cy - u * .48 + i * s * .21),
                  (cx, cy - u * .48 + i * s * .21)], color, lw * .85)
    elif name == 'star':
        shape(MSO_SHAPE.STAR_5_POINT, cx - u * .86, cy - u * .86,
              s * .86, s * .86, fill=color)


# ------------------------------------------------------------------ header ---
textbox(0.65, 0.50, 12.03, 0.55,
        [[('Re-energizing our GTM — ', 28, True, BLACK),
          ('three plays, one motion', 28, True, PURPLE)]])
textbox(0.65, 1.12, 12.03, 0.32,
        [('Industry-led, AI-powered, ecosystem-enabled GTM to accelerate '
          'platform adoption.', 16, False, GREY_1)])

# ------------------------------------------------------------ pillar cards ---
TOP, CH, CW = 1.62, 3.24, 3.78
PX = [0.65, 4.775, 8.90]

LIST_PILLARS = {
    0: [('target', 'Deepen vertical alignment'),
        ('people', 'CXO access & engagement'),
        ('chart_up', 'Activate BPO platform conversations')],
    2: [('cloud', 'Hyperscalers'),
        ('arrow_up', 'Startup plays — Krisp.ai, Stutt.ai'),
        ('nodes', 'Joint GTM initiatives')],
}
HEADS = [('01', 'VERTICAL SALES', 'Drive demand', PURPLE_DEEP),
         ('02', 'DBS + DPO', 'Unify the story', PURPLE),
         ('03', 'ECOSYSTEM INNOVATION', 'Scale the reach', BLUE)]

for i, (num, head, sub, col) in enumerate(HEADS):
    x = PX[i]
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, TOP, CW, CH,
          fill=WHITE, line=GREY_3, adj=0.03)
    chip(MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.26, TOP + 0.24, 0.52, 0.44,
         [[(num, 15, True, WHITE)]], fill=col, adj=0.2)
    textbox(x + 0.92, TOP + 0.22, CW - 1.18, 0.24, [(head, 12.5, True, col)])
    textbox(x + 0.92, TOP + 0.50, CW - 1.18, 0.24, [(sub, 12, False, BLACK)])

    if i in LIST_PILLARS:
        for j, (icon, text) in enumerate(LIST_PILLARS[i]):
            ry = TOP + 1.00 + j * 0.70
            shape(MSO_SHAPE.OVAL, x + 0.62, ry + 0.04, 0.44, 0.44, fill=GREY_4)
            glyph(icon, x + 0.84, ry + 0.26, 0.24, col)
            textbox(x + 1.22, ry, 2.30, 0.52, [(text, 10, False, BLACK)],
                    anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)
    else:
        for sx, lab, g in ((0.35, 'DBS', 'database'), (2.08, 'DPO', 'cubes')):
            shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + sx, TOP + 1.02, 1.35, 1.05,
                  fill=col, adj=0.11)
            glyph(g, x + sx + 0.675, TOP + 1.38, 0.32, WHITE)
            textbox(x + sx, TOP + 1.66, 1.35, 0.28, [(lab, 12, True, WHITE)],
                    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        chip(MSO_SHAPE.OVAL, x + 1.71, TOP + 1.36, 0.36, 0.36,
             [[('+', 14, True, col)]], fill=WHITE, line=col)
        poly([(x + CW / 2, TOP + 2.14), (x + CW / 2, TOP + 2.30)], col, 1.1)
        poly([(x + CW / 2 - 0.06, TOP + 2.24), (x + CW / 2, TOP + 2.30),
              (x + CW / 2 + 0.06, TOP + 2.24)], col, 1.1)
        chip(MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.26, TOP + 2.38, CW - 0.52, 0.66,
             [[('ONE INTEGRATED NARRATIVE', 10.5, True, col)],
              [('Greater customer value together', 9.5, False, BLACK)]],
             fill=GREY_4, adj=0.14)

    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.26, TOP + CH - 0.16, CW - 0.52, 0.06,
          fill=col, adj=0.5)

for gx in (4.60, 8.725):
    shape(MSO_SHAPE.ISOSCELES_TRIANGLE, gx - 0.10, TOP + 1.50, 0.20, 0.26,
          fill=GREY_3, rot=90)

# ------------------------------------------------------- achievement banner ---
BY, BH = 5.02, 0.82
shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, BY, 12.03, BH, fill=ICE_BLUE, adj=0.16)
BX = 6.665 - (0.56 + 0.36 + 0.24 + 4.40) / 2      # centre the banner group
shape(MSO_SHAPE.OVAL, BX, BY + 0.13, 0.56, 0.56, fill=PURPLE)
glyph('star', BX + 0.28, BY + 0.41, 0.30, WHITE)
poly([(BX + 0.86, BY + 0.18), (BX + 0.86, BY + 0.64)], GREY_3, 0.75)
textbox(BX + 1.10, BY, 4.45, BH,
        [[('1st DPO ADVISOR DAY COMPLETED', 15, True, PURPLE_DEEP)],
         [('A strong start to our platform-led GTM motion', 11, False, GREY_1)]],
        anchor=MSO_ANCHOR.MIDDLE, space_after=2)

# --------------------------------------------------------------- metric row ---
TW, TGAP, TY, TH = 2.82, 0.25, 6.00, 0.98
METRICS = [('chart_up', '89%', 'Willing to take DPO\nplatform to market'),
           ('people', '20+', 'Sales enablement\nsessions'),
           ('target', '5+', 'CFO roundtables\nplanned, FY27'),
           ('doc', '15+', 'Analyst briefings\nplanned, FY27')]
for i, (icon, num, label) in enumerate(METRICS):
    x = 0.65 + i * (TW + TGAP)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, TY, TW, TH,
          fill=WHITE, line=GREY_3, adj=0.08)
    shape(MSO_SHAPE.OVAL, x + 0.18, TY + 0.23, 0.52, 0.52, fill=PURPLE)
    glyph(icon, x + 0.44, TY + 0.49, 0.28, WHITE)
    textbox(x + 0.84, TY + 0.10, TW - 1.00, 0.38, [(num, 21, True, PURPLE)],
            anchor=MSO_ANCHOR.MIDDLE)
    textbox(x + 0.84, TY + 0.50, TW - 1.00, 0.40,
            [[(ln, 9.5, False, BLACK)] for ln in label.split('\n')],
            line_spacing=1.05)

prs.save(OUT)
print('saved', OUT, '| shapes:', len(slide.shapes))
