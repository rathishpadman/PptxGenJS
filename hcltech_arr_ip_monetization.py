"""HCLTech slide: ARR expansion through IP monetization."""
import sys
sys.path.insert(0, '.')
from pptx import Presentation
from hcltech_board_common import *          # noqa: F401,F403
import hcltech_board_common

TEMPLATE = '/root/.claude/skills/synced/hcltech-deck/assets/hcltech_base.pptx'
OUT = '/home/user/PptxGenJS/hcltech_arr_ip_monetization.pptx'

prs = Presentation(TEMPLATE)
slide = prs.slides.add_slide(prs.slide_layouts[16])
use(slide.shapes)

textbox(0.65, 0.50, 12.03, 0.55,
        [[('ARR grows 22–25% by FY30, ', 28, True, BLACK),
          ('driven by IP monetization', 28, True, PURPLE)]])
textbox(0.65, 1.12, 12.03, 0.32,
        [('Resource ARR holds as a stable base while the Platform IP layer '
          'compounds from FY27.', 16, False, GREY_1)])

# ---- compounded index: 1.025 x 1.04 x 1.06 x 1.08 = 1.220 -------------------
RATES = [('FY27', 0.025, '+2.5%'), ('FY28', 0.040, '+4.0%'),
         ('FY29', 0.060, '+6.0%'), ('FY30', 0.080, '+8.0%')]
idx, steps = 100.0, []
for lab, r, txt in RATES:
    delta = idx * r
    steps.append((lab, idx, idx + delta, txt))
    idx += delta
TOTAL_GROWTH = idx - 100.0                    # 22.04 index points

BASE_Y, BASE_H = 6.30, 1.70                   # FY26 slab (compressed)
BASE_TOP = BASE_Y - BASE_H
GROWTH_H = 2.09                               # the whole growth layer
PPU = GROWTH_H / TOTAL_GROWTH                 # inches per index point

X0, COLW, BW = 1.15, 1.525, 1.05
cx = [X0 + 0.76 + i * COLW for i in range(6)]

rule(X0, BASE_Y, 9.15, GREY_3, 1.0)

# FY26 base
shape(MSO_SHAPE.RECTANGLE, cx[0] - BW / 2, BASE_TOP, BW, BASE_H, fill=PURPLE_DEEP)

# floating increments
prev_top = BASE_TOP
for i, (lab, lo, hi, txt) in enumerate(steps):
    top = BASE_TOP - (hi - 100.0) * PPU
    bot = BASE_TOP - (lo - 100.0) * PPU
    shape(MSO_SHAPE.RECTANGLE, cx[i + 1] - BW / 2, top, BW, bot - top,
          fill=LAVENDER)
    textbox(cx[i + 1] - BW / 2, top - 0.30, BW, 0.26, [(txt, 13, True, PURPLE)],
            align=PP_ALIGN.CENTER)
    rule(cx[i] + BW / 2, bot, COLW - BW, GREY_2, 0.75, dash=True)
    prev_top = top

# FY30 total: base plus the accumulated IP layer
tx = cx[5] - BW / 2
shape(MSO_SHAPE.RECTANGLE, tx, BASE_TOP, BW, BASE_H, fill=PURPLE_DEEP)
shape(MSO_SHAPE.RECTANGLE, tx, BASE_TOP - GROWTH_H, BW, GROWTH_H, fill=PURPLE)
rule(cx[4] + BW / 2, prev_top, COLW - BW, GREY_2, 0.75, dash=True)

for i, lab in enumerate(['FY26', 'FY27', 'FY28', 'FY29', 'FY30']):
    textbox(cx[i] - BW / 2, BASE_Y + 0.14, BW, 0.26, [(lab, 12, True, BLACK)],
            align=PP_ALIGN.CENTER)
textbox(cx[5] - BW / 2, BASE_Y + 0.14, BW, 0.52,
        [[('FY30', 12, True, BLACK)], [('TOTAL', 12, True, BLACK)]],
        align=PP_ALIGN.CENTER, line_spacing=1.0)

textbox(X0, 6.98, 8.60, 0.22,
        [('Indexed to FY26 = 100. Year-on-year growth compounds to +22%; the '
          'growth layer is scaled for legibility.', 8.5, False, GREY_2)])

# ---- span bracket over the whole build ------------------------------------
LY = 1.88
poly([(1.52, BASE_TOP - 0.18), (1.52, LY), (9.80, LY), (9.80, 2.40)],
     PURPLE, 1.25)
arrowhead(9.80, 2.46, PURPLE, 0.075)
chip(MSO_SHAPE.ROUNDED_RECTANGLE, 5.02, 1.56, 3.30, 0.64,
     [[('+22% – 25%', 17, True, PURPLE)],
      [('INCREASE (FY26 TO FY30)', 9, True, GREY_1)]],
     fill=WHITE, line=PURPLE, adj=0.5)

# ---- right-hand legend brackets -------------------------------------------
BX = 10.34
for y0, y1, icon, name, desc, col in (
        (BASE_TOP - GROWTH_H, BASE_TOP, 'cube', 'Platform IP Revenue',
         'New revenue layer\nfrom IP monetization', PURPLE),
        (BASE_TOP, BASE_Y, 'database', 'Resource ARR',
         'Stable base from\nexisting operations', PURPLE_DEEP)):
    vrule(BX, y0, y1 - y0, col, 1.0)
    rule(BX, y0, 0.10, col, 1.0)
    rule(BX, y1, 0.10, col, 1.0)
    mid = (y0 + y1) / 2
    shape(MSO_SHAPE.OVAL, BX + 0.20, mid - 0.24, 0.48, 0.48, fill=col)
    glyph(icon, BX + 0.44, mid, 0.26, WHITE)
    textbox(BX + 0.80, mid - 0.42, 1.84, 0.24, [(name, 11, True, col)])
    textbox(BX + 0.80, mid - 0.14, 1.84, 0.56,
            [[(ln, 9.5, False, GREY_1)] for ln in desc.split('\n')],
            line_spacing=1.1)

prs.save(OUT)
print('saved', OUT, '| shapes:', len(slide.shapes),
      '| total index %.1f' % idx)
