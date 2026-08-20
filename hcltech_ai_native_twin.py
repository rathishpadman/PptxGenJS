"""HCLTech slide: Shifting from People to Platforms — The AI Native Twin Strategy."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TEMPLATE = '/root/.claude/skills/synced/hcltech-deck/assets/hcltech_base.pptx'
OUT = '/home/user/PptxGenJS/hcltech_ai_native_twin.pptx'

BLACK       = RGBColor(0x00, 0x00, 0x00)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
BLUE        = RGBColor(0x0F, 0x5F, 0xDC)
BRIGHT_BLUE = RGBColor(0x3C, 0x91, 0xFF)
LIGHT_BLUE  = RGBColor(0x8C, 0xC8, 0xFA)
ICE_BLUE    = RGBColor(0xDC, 0xE6, 0xF0)
GREY_1      = RGBColor(0x82, 0x91, 0xA0)
GREY_3      = RGBColor(0xC8, 0xD2, 0xDD)
GREY_4      = RGBColor(0xE6, 0xEB, 0xF5)
FONT = 'Aptos'

prs = Presentation(TEMPLATE)
slide = prs.slides.add_slide(prs.slide_layouts[17])
shapes = slide.shapes
I = Inches


def style(run, size, *, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def fill_tf(tf, paras, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            line_spacing=None, space_after=0):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, prun in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for text, size, bold, color in prun:
            style(p.add_run(), size, bold=bold, color=color)
            p.runs[-1].text = text


def textbox(x, y, w, h, runs, **kw):
    tb = shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, runs if runs and isinstance(runs[0], list) else [runs], **kw)
    return tb


def shape(kind, x, y, w, h, *, fill=None, line=None, line_w=0.75, adj=None, rot=None):
    s = shapes.add_shape(kind, I(x), I(y), I(w), I(h))
    if adj is not None:
        vals = adj if isinstance(adj, (list, tuple)) else [adj]
        for i, v in enumerate(vals):
            if i < len(s.adjustments):
                s.adjustments[i] = v
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    if rot is not None:
        s.rotation = rot
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    return s


def polyline(pts, color, width, close=False):
    b = shapes.build_freeform(I(pts[0][0]), I(pts[0][1]))
    b.add_line_segments([(I(x), I(y)) for x, y in pts[1:]], close=close)
    s = b.convert_to_shape()
    s.fill.background()
    s.line.color.rgb = color
    s.line.width = Pt(width)
    s.shadow.inherit = False
    return s


# ---------------------------------------------------------------- eyebrow ----
pill = shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, 0.42, 2.48, 0.28,
             fill=ICE_BLUE, adj=0.5)
pill.text_frame.margin_left = pill.text_frame.margin_right = 0
pill.text_frame.margin_top = pill.text_frame.margin_bottom = 0
fill_tf(pill.text_frame, [[('STRATEGIC PLATFORM DELIVERY', 9, True, BLUE)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ------------------------------------------------------- title + subtitle ----
title = slide.shapes.title
title.left, title.top, title.width, title.height = I(0.65), I(0.78), I(12.03), I(1.08)
fill_tf(title.text_frame,
        [[('Shifting from People to Platforms:', 28, True, BLACK)],
         [('The AI Native Twin Strategy', 28, True, BLUE)]],
        line_spacing=0.95)

sub = [ph for ph in slide.placeholders if ph.placeholder_format.idx == 12][0]
sub._element.getparent().remove(sub._element)
textbox(0.66, 1.94, 12.00, 0.56,
        [('Bypassing organizational resistance by incubating standalone, AI-driven '
          'operations on the edge to smoothly transition legacy workflows.',
          16, False, GREY_1)], line_spacing=1.1)

# ------------------------------------------------------------- phase cards ---
PX = [0.65, 4.775, 8.90]
PW, TOP, CH = 3.78, 2.62, 3.15
PAD = 0.30

CARDS = [
    ('PHASE 01', '1. Edge Incubation',
     'Built in isolation — past mainstream resistance and organizational antibodies.',
     'Isolated Edge Blueprint', 'nodes'),
    ('PHASE 02', '2. Targeted Maturity',
     'Deep agents trained fast, on a minimal footprint.',
     '2 Domain Experts + 2 AI Engineers', 'target'),
    ('PHASE 03', '3. Phased Cutover',
     'By department, by geography — never a single flip of the button.',
     'Low-Risk Incremental Rollout', 'handover'),
]


def icon(kind, tx, ty, size):
    """Draw a glyph inside a tile whose top-left is (tx, ty)."""
    cx, cy = tx + size / 2, ty + size / 2
    if kind == 'nodes':
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.065, cy - 0.19, 0.13, 0.13,
              fill=BLUE, adj=0.25)
        for dx in (-0.20, -0.05, 0.10):
            shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + dx, cy + 0.075, 0.10, 0.10,
                  fill=BLUE, adj=0.25)
        polyline([(cx - 0.15, cy + 0.025), (cx + 0.15, cy + 0.025)], BLUE, 1.25)
    elif kind == 'target':
        shape(MSO_SHAPE.OVAL, cx - 0.175, cy - 0.175, 0.35, 0.35, fill=None,
              line=BLUE, line_w=1.5)
        shape(MSO_SHAPE.OVAL, cx - 0.10, cy - 0.10, 0.20, 0.20, fill=None,
              line=BLUE, line_w=1.5)
        shape(MSO_SHAPE.OVAL, cx - 0.042, cy - 0.042, 0.085, 0.085, fill=BLUE)
    elif kind == 'handover':
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.25, cy - 0.065, 0.13, 0.13,
              fill=BLUE, adj=0.25)
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + 0.12, cy - 0.065, 0.13, 0.13,
              fill=None, line=BLUE, line_w=1.5, adj=0.25)
        shape(MSO_SHAPE.RIGHT_ARROW, cx - 0.06, cy - 0.037, 0.12, 0.075,
              fill=BLUE, adj=(0.5, 0.55))


for i, (phase, head, body, chip, glyph) in enumerate(CARDS):
    x = PX[i]
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, TOP, PW, CH,
          fill=WHITE, line=GREY_3, adj=0.03)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + PAD, TOP - 0.03, PW - 2 * PAD, 0.06,
          fill=BLUE, adj=0.5)

    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + PAD, TOP + 0.26, 0.62, 0.62,
          fill=ICE_BLUE, adj=0.24)
    icon(glyph, x + PAD, TOP + 0.26, 0.62)

    ph = shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + PW - PAD - 1.02, TOP + 0.40,
               1.02, 0.30, fill=ICE_BLUE, adj=0.5)
    ph.text_frame.margin_left = ph.text_frame.margin_right = 0
    ph.text_frame.margin_top = ph.text_frame.margin_bottom = 0
    fill_tf(ph.text_frame, [[(phase, 8.5, True, BLUE)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    textbox(x + PAD, TOP + 1.06, PW - 2 * PAD, 0.34,
            [(head, 16, True, BLACK)], anchor=MSO_ANCHOR.MIDDLE)
    textbox(x + PAD, TOP + 1.48, PW - 2 * PAD, 0.80,
            [(body, 11, False, GREY_1)], line_spacing=1.15)

    cy = TOP + 2.40
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + PAD, cy, PW - 2 * PAD, 0.44,
          fill=GREY_4, adj=0.12)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + PAD, cy, 0.06, 0.44, fill=BLUE, adj=0.5)
    textbox(x + PAD + 0.22, cy, PW - 2 * PAD - 0.34, 0.44,
            [(chip, 9.5, True, BLACK)], anchor=MSO_ANCHOR.MIDDLE)

# ------------------------------------------------------------- flow arrows ---
for gx in (4.60, 8.725):
    gy = TOP + 1.45
    shape(MSO_SHAPE.OVAL, gx - 0.17, gy - 0.17, 0.34, 0.34,
          fill=WHITE, line=GREY_3)
    polyline([(gx - 0.045, gy - 0.075), (gx + 0.048, gy), (gx - 0.045, gy + 0.075)],
             BLUE, 1.5)

# ------------------------------------------------------------- impact band ---
band = shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, 6.05, 12.03, 0.82, adj=0.10)
band.fill.gradient()
band.fill.gradient_stops[0].color.rgb = BLUE
band.fill.gradient_stops[0].position = 0.0
band.fill.gradient_stops[1].color.rgb = BRIGHT_BLUE
band.fill.gradient_stops[1].position = 1.0
band.fill.gradient_angle = 0.0

shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.95, 6.20, 0.52, 0.52,
      fill=BRIGHT_BLUE, adj=0.26)
shape(MSO_SHAPE.UP_ARROW, 1.08, 6.33, 0.26, 0.26, fill=WHITE,
      adj=(0.42, 0.42), rot=45)

textbox(1.68, 6.05, 10.65, 0.82,
        [[('MEASURABLE BUSINESS IMPACT', 8.5, True, LIGHT_BLUE)],
         [('Radical Efficiency: ', 12.5, True, WHITE),
          ('300-person workflows into hyper-productive native twins — '
           '6 months from architectural approval.', 12.5, False, WHITE)]],
        anchor=MSO_ANCHOR.MIDDLE, space_after=2)

prs.save(OUT)
print('saved', OUT, '| shapes:', len(slide.shapes))
