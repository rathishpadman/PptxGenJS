"""HCLTech board slide: engineered vs conventional agents (anomaly detection)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TEMPLATE = '/root/.claude/skills/synced/hcltech-deck/assets/hcltech_base.pptx'
OUT = '/home/user/PptxGenJS/hcltech_engineered_agents.pptx'

BLACK       = RGBColor(0x00, 0x00, 0x00)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
PURPLE      = RGBColor(0x5F, 0x1E, 0xBE)
PURPLE_DEEP = RGBColor(0x41, 0x14, 0x82)
BLUE        = RGBColor(0x0F, 0x5F, 0xDC)
ICE_BLUE    = RGBColor(0xDC, 0xE6, 0xF0)
GREY_1      = RGBColor(0x82, 0x91, 0xA0)
GREY_3      = RGBColor(0xC8, 0xD2, 0xDD)
GREY_4      = RGBColor(0xE6, 0xEB, 0xF5)
FONT = 'Aptos'

prs = Presentation(TEMPLATE)
slide = prs.slides.add_slide(prs.slide_layouts[16])
S = slide.shapes
I = Inches


def style(run, size, *, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def fill_tf(tf, paras, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            line_spacing=None, space_after=0):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, prun in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for text, size, bold, color in prun:
            style(p.add_run(), size, bold=bold, color=color)
            p.runs[-1].text = text


def textbox(x, y, w, h, runs, **kw):
    tb = S.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, runs if runs and isinstance(runs[0], list) else [runs], **kw)
    return tb


def shape(kind, x, y, w, h, *, fill=None, line=None, line_w=0.75, adj=None):
    s = S.add_shape(kind, I(x), I(y), I(w), I(h))
    if adj is not None:
        vals = adj if isinstance(adj, (list, tuple)) else [adj]
        for i, v in enumerate(vals):
            if i < len(s.adjustments):
                s.adjustments[i] = v
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    return s


def poly(pts, color, width):
    b = S.build_freeform(I(pts[0][0]), I(pts[0][1]))
    b.add_line_segments([(I(x), I(y)) for x, y in pts[1:]], close=False)
    s = b.convert_to_shape()
    s.fill.background()
    s.line.color.rgb = color
    s.line.width = Pt(width)
    s.shadow.inherit = False
    return s


def chip(x, y, w, h, lines, *, fill, line=None, adj=0.5, pad=0.06):
    c = shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line, adj=adj)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = I(pad)
    tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, lines, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
            line_spacing=0.95)
    return c


def arrow(cx, cy, color=GREY_1):
    poly([(cx - 0.08, cy), (cx + 0.08, cy)], color, 1.0)
    poly([(cx + 0.03, cy - 0.055), (cx + 0.08, cy), (cx + 0.03, cy + 0.055)],
         color, 1.0)


# ------------------------------------------------------------------ header ---
textbox(0.65, 0.50, 12.03, 0.55,
        [[('Engineered agents: ', 28, True, BLACK),
          ('95% accuracy, 70% fewer tokens', 28, True, PURPLE)]])
textbox(0.65, 1.12, 12.03, 0.32,
        [('Anomaly detection — retrieval, rules and validation before reasoning, '
          'not a bigger prompt.', 16, False, GREY_1)])

# --------------------------------------------------- layer 1: three headline --
TW, GAP = 3.85, 0.24
PRIMARY = [
    ('ACCURACY', '60%', '95%', '+35 pts'),
    ('FALSE POSITIVE RISK', '40%', '10%', '75% lower'),
    ('INPUT TOKENS', '12,685', '3,735', '70% lower'),
]
for i, (label, before, after, delta) in enumerate(PRIMARY):
    x = 0.65 + i * (TW + GAP)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.60, TW, 1.55,
          fill=WHITE, line=GREY_3, adj=0.03)
    textbox(x, 1.78, TW, 0.22, [(label, 9.5, True, GREY_1)], align=PP_ALIGN.CENTER)
    textbox(x, 2.06, TW, 0.60,
            [[(before, 26, False, GREY_1), ('  →  ', 18, False, GREY_3),
              (after, 30, True, PURPLE)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    chip(x + TW / 2 - 0.85, 2.72, 1.70, 0.30,
         [[(delta, 10, True, PURPLE)]], fill=ICE_BLUE)

# ------------------------------------------- layer 2: two supporting metrics --
shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, 3.28, 12.03, 0.62,
      fill=GREY_4, adj=0.16)
poly([(6.665, 3.40), (6.665, 3.78)], GREY_3, 0.75)
SECONDARY = [
    ('LATENCY', '12.0s', '10.9s', '9% faster'),
    ('DECISION QUALITY', 'Partial', 'Validated', 'Higher confidence'),
]
for i, (label, before, after, delta) in enumerate(SECONDARY):
    textbox(0.65 + i * 6.015, 3.28, 6.015, 0.62,
            [[(label + '    ', 9, True, GREY_1),
              (before, 13, False, GREY_1), ('  →  ', 11, False, GREY_3),
              (after, 13, True, PURPLE), ('    ' + delta, 9.5, False, PURPLE)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

textbox(0.65, 3.98, 12.03, 0.20,
        [('Anomaly detection benchmark; single use case.', 8.5, False, GREY_1)])

# ------------------------------------------------ architecture: same baseline -
textbox(0.65, 4.34, 12.03, 0.22,
        [[('Same input. ', 10.5, False, GREY_1),
          ('Four more steps before the model reasons.', 10.5, True, PURPLE)]])

CX0, CW_, CGAP = 2.35, 1.54, 0.22
ROWS = [
    (4.66, 'CONVENTIONAL', GREY_1,
     [('Input', False), ('LLM Analysis', False), ('Response', False)]),
    (5.62, 'ENGINEERED', PURPLE,
     [('Input', False), ('Context Retrieval', True), ('Rules & Policy', True),
      ('Historical Comparison', True), ('Validation Layer', True),
      ('Decision & Action', False)]),
]
for ry, rlabel, rcolor, steps in ROWS:
    textbox(0.65, ry, 1.55, 0.52, [(rlabel, 10, True, rcolor)],
            anchor=MSO_ANCHOR.MIDDLE)
    for j, (name, added) in enumerate(steps):
        cx = CX0 + j * (CW_ + CGAP)
        if added:
            chip(cx, ry, CW_, 0.52, [[(name, 9.5, True, WHITE)]],
                 fill=PURPLE, adj=0.14)
        else:
            chip(cx, ry, CW_, 0.52,
                 [[(name, 9.5, True, BLACK if rcolor is GREY_1 else PURPLE)]],
                 fill=GREY_4 if rcolor is GREY_1 else ICE_BLUE, adj=0.14)
        if j < len(steps) - 1:
            arrow(cx + CW_ + CGAP / 2, ry + 0.26, GREY_3)

# ------------------------------------------------------------- closing band ---
band = shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, 6.42, 12.03, 0.56, adj=0.14)
band.fill.gradient()
band.fill.gradient_stops[0].color.rgb = PURPLE_DEEP
band.fill.gradient_stops[0].position = 0.0
band.fill.gradient_stops[1].color.rgb = PURPLE
band.fill.gradient_stops[1].position = 1.0
band.fill.gradient_angle = 0.0
textbox(0.65, 6.42, 12.03, 0.56,
        [[('Context, rules and validation before reasoning — ', 12, False, WHITE),
          ('better decisions on a third of the compute.', 12, True, WHITE)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print('saved', OUT, '| shapes:', len(slide.shapes))
