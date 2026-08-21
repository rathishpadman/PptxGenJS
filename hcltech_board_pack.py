"""HCLTech board pack — five slides."""
import sys
sys.path.insert(0, '.')
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from hcltech_board_common import *          # noqa: F401,F403
from hcltech_board_common import _S         # noqa: F401

TEMPLATE = '/root/.claude/skills/synced/hcltech-deck/assets/hcltech_base.pptx'
OUT = '/home/user/PptxGenJS/hcltech_board_pack.pptx'

prs = Presentation(TEMPLATE)


def new_slide():
    s = prs.slides.add_slide(prs.slide_layouts[16])
    use(s.shapes)
    return s


def header(title_runs, subtitle):
    textbox(0.65, 0.50, 12.03, 0.55, [[(t, 28, True, c) for t, c in title_runs]])
    textbox(0.65, 1.12, 12.03, 0.32, [(subtitle, 16, False, GREY_1)])


# ═══════════════════════════════════ 1 · Restructuring for Autonomous Ops ════
def slide_restructuring():
    new_slide()
    header([('Restructuring for ', BLACK), ('Autonomous Operations', PURPLE)],
           'Four specialized teams aligned to create, deploy and run one '
           'AI-native platform.')

    GROUPS = [
        (0.65, 5.30, PURPLE_DEEP, 'code', 'CREATE THE PLATFORM',
         'Product + AI Engineering',
         [('cube', 'Product Team', '80+', 'Define platform\nand outcomes'),
          ('nodes', 'AI Engineering', '250+', 'Build and improve\nthe platform')]),
        (6.25, 3.05, PURPLE, 'plane', 'DEPLOY THE PLATFORM',
         'Forward Deployed Engineers',
         [('people', 'FDEs', '750+', 'Deploy and optimize\nfor clients')]),
        (9.60, 3.08, BLUE, 'shield', 'RUN THE PLATFORM', 'Operations',
         [('gear', 'Operations', '35,000+', 'Operate and govern\nat scale')]),
    ]
    GT, GH = 1.98, 2.72
    for gx, gw, col, icon, head, sub, cards in GROUPS:
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, gx, GT, gw, GH,
              line=col, line_w=1.0, adj=0.045)
        # header chip notches over the container's top edge
        hx, hw = gx + 0.20, gw - 0.40
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, hx, GT - 0.28, hw, 0.56,
              fill=WHITE, line=col, line_w=1.0, adj=0.16)
        shape(MSO_SHAPE.OVAL, hx + 0.14, GT - 0.11, 0.34, 0.34,
              fill=WHITE, line=col, line_w=1.0)
        glyph(icon, hx + 0.31, GT + 0.06, 0.18, col)
        textbox(hx + 0.58, GT - 0.23, hw - 0.68, 0.23, [(head, 10, True, col)])
        textbox(hx + 0.58, GT + 0.02, hw - 0.68, 0.22, [(sub, 9, False, BLACK)])

        n = len(cards)
        inner = gw - 0.48
        cw = (inner - 0.20 * (n - 1)) / n
        for j, (cicon, cname, cnum, ccap) in enumerate(cards):
            cx = gx + 0.24 + j * (cw + 0.20)
            ct, ch = GT + 0.44, 2.14
            shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, ct, cw, ch,
                  fill=WHITE, line=GREY_3, adj=0.05)
            shape(MSO_SHAPE.OVAL, cx + cw / 2 - 0.31, ct + 0.16, 0.62, 0.62, fill=col)
            glyph(cicon, cx + cw / 2, ct + 0.47, 0.32, WHITE)
            textbox(cx, ct + 0.90, cw, 0.28, [(cname, 13, True, col)],
                    align=PP_ALIGN.CENTER)
            rule(cx + 0.34, ct + 1.24, cw - 0.68, GREY_4)
            textbox(cx, ct + 1.32, cw, 0.42, [(cnum, 24, True, col)],
                    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            textbox(cx + 0.10, ct + 1.76, cw - 0.20, 0.36,
                    [[(ln, 10, True, BLACK)] for ln in ccap.split('\n')],
                    align=PP_ALIGN.CENTER, line_spacing=1.05)

        gcx = gx + gw / 2
        shape(MSO_SHAPE.OVAL, gcx - 0.055, GT + GH - 0.055, 0.11, 0.11, fill=col)
        vrule(gcx, GT + GH + 0.06, 0.72, col, dash=True)
        arrowhead(gcx, GT + GH + 0.82, col)

    band = shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, 5.58, 12.03, 0.92, adj=0.10)
    gradient(band, PURPLE_DEEP, BLUE)
    shape(MSO_SHAPE.OVAL, 3.59, 5.76, 0.56, 0.56, fill=WHITE)
    glyph('layers', 3.87, 6.04, 0.30, PURPLE)
    textbox(4.35, 5.58, 5.40, 0.92,
            [[('AUTONOMOUS OPERATIONS PLATFORM', 17, True, WHITE)],
             [('AI-native   |   Scalable   |   Secure   |   Reusable',
               11, False, PERIWINKLE)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=3)

    textbox(0.65, 6.70, 12.03, 0.28,
            [('One Platform.  Four Specialized Teams.  End-to-End Value.',
              12.5, True, PURPLE_DEEP)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    rule(1.05, 6.84, 2.35, GREY_3)
    rule(9.93, 6.84, 2.35, GREY_3)


# ═══════════════════════════════════════════════ 2 · FY30 Aspiration ═════════
def slide_fy30():
    new_slide()
    header([('FY30 aspiration: ', BLACK), ('$1.5B through engineered growth', PURPLE)],
           '2x revenue, 4x business growth, and 430 bps of EBIT expansion.')

    chip(MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, 1.66, 7.55, 0.46,
         [[('2x Revenue  =  4x Business Growth', 14, True, BLUE)]],
         fill=ICE_BLUE, adj=0.2)

    BASE, TOP_V = 6.06, 1500.0
    PLOT_H = 3.40                       # 1500M spans this many inches
    def h_of(v):
        return v / TOP_V * PLOT_H

    rule(0.85, BASE, 7.55, GREY_3, 1.0)
    BW = 1.20
    for cx, lab in ((1.95, 'FY26'), (4.25, 'FY27-30'), (6.55, 'FY30')):
        chip(MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.62, BASE + 0.12, 1.24, 0.32,
             [[(lab, 11, True, BLACK)]], fill=WHITE, line=GREY_3, adj=0.16)

    # FY26
    y = BASE - h_of(714)
    shape(MSO_SHAPE.RECTANGLE, 1.95 - BW / 2, y, BW, h_of(714), fill=ICE_BLUE)
    textbox(1.95 - BW / 2, y - 0.30, BW, 0.26, [('714 M', 12.5, True, BLACK)],
            align=PP_ALIGN.CENTER)

    # FY27-30: retained base with the at-risk block hatched above it
    yb = BASE - h_of(400)
    shape(MSO_SHAPE.RECTANGLE, 4.25 - BW / 2, yb, BW, h_of(400), fill=ICE_BLUE)
    textbox(4.25 - BW / 2, yb + 0.10, BW, 0.52,
            [[('~375 -', 11.5, True, BLACK)], [('400 M', 11.5, True, BLACK)]],
            align=PP_ALIGN.CENTER, line_spacing=1.0)
    yr = BASE - h_of(714)
    hatch(4.25 - BW / 2, yr, BW, h_of(314), GREY_2, WHITE, line=GREY_2)
    textbox(4.25 - BW / 2, yr + 0.16, BW, 0.24, [('~Rev at Risk', 10, True, BLACK)],
            align=PP_ALIGN.CENTER)
    bx = 4.25 - BW / 2 - 0.12
    vrule(bx, yr, h_of(314), GREY_2, 1.0)
    rule(bx, yr, 0.10, GREY_2, 1.0)
    rule(bx, yr + h_of(314), 0.10, GREY_2, 1.0)
    textbox(1.98, yr + 0.02, 1.38, 0.46,
            [[('AI led', 9.5, False, GREY_1)], [('Compression', 9.5, False, GREY_1)]],
            align=PP_ALIGN.RIGHT, line_spacing=1.0)
    textbox(1.98, yr + h_of(314) - 0.24, 1.38, 0.24,
            [('Depletion', 9.5, False, GREY_1)], align=PP_ALIGN.RIGHT)

    # FY30
    yb2 = BASE - h_of(400)
    shape(MSO_SHAPE.RECTANGLE, 6.55 - BW / 2, yb2, BW, h_of(400), fill=ICE_BLUE)
    textbox(6.55 - BW / 2, yb2 + 0.10, BW, 0.52,
            [[('~375 -', 11.5, True, BLACK)], [('400 M', 11.5, True, BLACK)]],
            align=PP_ALIGN.CENTER, line_spacing=1.0)
    y2 = BASE - h_of(1500)
    shape(MSO_SHAPE.RECTANGLE, 6.55 - BW / 2, y2, BW, h_of(1100), fill=LIGHT_BLUE)
    textbox(6.55 - BW / 2, y2 + h_of(1100) / 2 - 0.16, BW, 0.30,
            [('~1.1 B', 13, True, BLACK)], align=PP_ALIGN.CENTER)
    textbox(6.55 - BW / 2, y2 - 0.32, BW, 0.28, [('1.5 B', 13.5, True, BLACK)],
            align=PP_ALIGN.CENTER)

    poly([(4.90, yb), (6.00, y2 + 0.06)], GREY_1, 1.25)
    arrowhead(6.02, y2 + 0.02, GREY_1, 0.075, 'up')
    chip(MSO_SHAPE.OVAL, 5.05, 3.38, 0.84, 0.46, [[('4X', 14, True, BLACK)]],
         fill=GREY_4)

    # EBIT panel
    PX_, PW_ = 8.62, 4.06
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, PX_, 1.66, PW_, 4.86,
          fill=WHITE, line=GREY_3, adj=0.04)
    chip(MSO_SHAPE.ROUNDED_RECTANGLE, PX_ + 0.24, 1.86, PW_ - 0.48, 0.62,
         [[('EBIT EXPANSION', 13, True, PURPLE)],
          [('FY26 → FY30', 10, False, GREY_1)]], fill=GREY_4, adj=0.18)
    textbox(PX_, 2.62, PW_, 0.80,
            [[('+430', 44, True, BLUE), (' bps', 15, True, BLUE)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    X0, X1 = PX_ + 0.50, PX_ + PW_ - 0.50
    YB, VMIN, VRANGE, PLOTH = 5.84, 12.5, 5.7, 1.48
    pts = [(X0, 13.3), ((X0 + X1) / 2, 15.4), (X1, 17.6)]
    xy = [(px, YB - (v - VMIN) / VRANGE * PLOTH) for px, v in pts]
    b = _shapes().build_freeform(I(xy[0][0]), I(YB))
    b.add_line_segments([(I(px), I(py)) for px, py in xy] + [(I(X1), I(YB))],
                        close=True)
    a = b.convert_to_shape()
    a.fill.solid(); a.fill.fore_color.rgb = ICE_BLUE
    a.line.fill.background(); a.shadow.inherit = False
    poly([(px, py) for px, py in xy], BLUE, 2.25)
    rule(X0 - 0.12, YB, (X1 - X0) + 0.24, GREY_3)
    for px, py in xy:
        shape(MSO_SHAPE.OVAL, px - 0.065, py - 0.065, 0.13, 0.13,
              fill=BLUE, line=WHITE, line_w=1.25)
    textbox(X0 - 0.04, xy[0][1] - 0.30, 0.70, 0.24, [('13.3%', 9.5, True, GREY_1)])
    textbox(X1 - 0.66, xy[2][1] - 0.32, 0.70, 0.24, [('17.6%', 10.5, True, BLUE)],
            align=PP_ALIGN.RIGHT)
    textbox(X0 - 0.04, YB + 0.10, 0.70, 0.22, [('FY26', 10, True, BLACK)])
    textbox(X1 - 0.66, YB + 0.10, 0.70, 0.22, [('FY30', 10, True, BLACK)],
            align=PP_ALIGN.RIGHT)

    shape(MSO_SHAPE.RIGHT_ARROW, 8.24, 3.84, 0.32, 0.42, fill=PURPLE_DEEP,
          adj=(0.55, 0.42))


def _shapes():
    import board_common
    return board_common._S


# ═══════════════════════════════════════════ 3 · Deliberate Growth ═══════════
def slide_growth():
    slide = new_slide()
    header([('Deliberate growth through ', BLACK),
            ('portfolio diversification', PURPLE)],
           '25 to 80 accounts by FY30 — broader base, deeper verticals, '
           'wider geographic spread.')

    PX = [0.65, 4.775, 8.90]
    PW, PT, PH = 3.78, 1.72, 4.35

    # ---- panel 1: account base
    x = PX[0]
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, PT, PW, PH, fill=WHITE, line=GREY_3,
          adj=0.03)
    textbox(x, PT + 0.26, PW, 0.34, [('80% of revenue', 16, True, BLACK)],
            align=PP_ALIGN.CENTER)
    textbox(x + 0.40, PT + 0.94, 1.20, 0.22, [('FY27', 10, False, GREY_1)],
            align=PP_ALIGN.CENTER)
    textbox(x + 2.18, PT + 0.94, 1.20, 0.22, [('FY30', 10, False, GREY_1)],
            align=PP_ALIGN.CENTER)
    textbox(x + 0.40, PT + 1.20, 1.20, 0.72, [('25', 42, True, BLACK)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(x + 1.62, PT + 1.20, 0.54, 0.72, [('→', 22, True, GREY_2)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(x + 2.18, PT + 1.20, 1.20, 0.72, [('80', 42, True, PURPLE)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for dx in (0.40, 2.18):
        textbox(x + dx, PT + 1.96, 1.20, 0.22, [('Accounts', 10, False, GREY_1)],
                align=PP_ALIGN.CENTER)
    rule(x + 0.42, PT + 2.60, PW - 0.84, GREY_3)
    textbox(x, PT + 2.86, PW, 0.46,
            [[('$24M', 21, True, BLACK), ('   →   ', 14, True, GREY_2),
              ('$15M', 21, True, PURPLE)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(x, PT + 3.38, PW, 0.22,
            [('Average revenue per account', 10, False, GREY_1)],
            align=PP_ALIGN.CENTER)

    # ---- panel 2: vertical mix table
    x = PX[1]
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, PT, PW, PH, fill=WHITE, line=GREY_3,
          adj=0.03)
    textbox(x, PT + 0.26, PW, 0.34, [('Vertical Mix', 16, True, BLACK)],
            align=PP_ALIGN.CENTER)
    COLS = [0.80, 0.58, 0.58, 1.42]
    tx, ty = x + 0.20, PT + 0.82
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, tx, ty, sum(COLS), 0.40, fill=PURPLE,
          adj=0.16)
    for j, (cw, lab) in enumerate(zip(COLS, ['', 'FY26\n($M)', 'FY30\n($M)',
                                             'Propositions'])):
        if not lab:
            continue
        cxx = tx + sum(COLS[:j])
        textbox(cxx, ty, cw, 0.40,
                [[(ln, 8.5, True, WHITE)] for ln in lab.split('\n')],
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.9)
    ROWS = [('TMT', '259', '500', 'Sales Ops, Customer Service'),
            ('FS', '113', '300', 'Insurance, Retail Banking'),
            ('Neo.AI', '106', '225', 'F&A, Procurement'),
            ('GEM', '65', '175', 'F&A, Procurement, HRO')]
    for i, row in enumerate(ROWS):
        ry = ty + 0.40 + i * 0.54
        if i % 2 == 0:
            shape(MSO_SHAPE.RECTANGLE, tx, ry, sum(COLS), 0.54, fill=GREY_4)
        for j, (cw, val) in enumerate(zip(COLS, row)):
            cxx = tx + sum(COLS[:j])
            textbox(cxx + 0.05, ry, cw - 0.10, 0.54,
                    [(val, 9.5 if j == 3 else 10.5, j == 0,
                      BLACK if j != 3 else GREY_1)],
                    align=PP_ALIGN.LEFT if j == 0 else (
                        PP_ALIGN.CENTER if j < 3 else PP_ALIGN.LEFT),
                    anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
    rule(tx, ty + 0.40 + len(ROWS) * 0.54, sum(COLS), GREY_3)

    # ---- panel 3: region mix donut
    x = PX[2]
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, PT, PW, PH, fill=WHITE, line=GREY_3,
          adj=0.03)
    textbox(x, PT + 0.26, PW, 0.34, [('Region Mix', 16, True, BLACK)],
            align=PP_ALIGN.CENTER)
    cd = CategoryChartData()
    cd.categories = ['Americas', 'EMEA', 'RoW']
    cd.add_series('FY30', (1000, 400, 100))
    gf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, I(x + 0.52), I(PT + 0.66),
                                I(PW - 1.04), I(2.42), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    plot = ch.plots[0]
    plot.vary_by_categories = True
    for pt, col in zip(plot.series[0].points, (PURPLE, LAVENDER, PERIWINKLE)):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = col
        pt.format.line.color.rgb = WHITE
        pt.format.line.width = Pt(1.5)
    for lbl, val, col in (('Americas  $1 Bn', 0, PURPLE),
                          ('EMEA  $400 Mn', 1, LAVENDER),
                          ('RoW  $100 Mn', 2, PERIWINKLE)):
        pass
    textbox(x + 0.30, PT + 3.22, PW - 0.60, 0.86,
            [[('Americas', 10.5, True, PURPLE), ('   $1 Bn', 10.5, False, BLACK)],
             [('EMEA', 10.5, True, LAVENDER), ('   $400 Mn', 10.5, False, BLACK)],
             [('RoW', 10.5, True, GREY_1), ('   $100 Mn', 10.5, False, BLACK)]],
            align=PP_ALIGN.CENTER, line_spacing=1.15)

    for gx in (4.60, 8.725):
        shape(MSO_SHAPE.ISOSCELES_TRIANGLE, gx - 0.10, PT + 1.90, 0.20, 0.26,
              fill=GREY_3, rot=90)


# ═════════════════════════════════════ 4 · Expanding Revenue (waterfall) ═════
def slide_arr():
    new_slide()
    header([('Expanding revenue ', BLACK), ('beyond traditional ARR', PURPLE)],
           'Transaction-based pricing adds a Platform IP layer on top of '
           'Resource ARR.')

    BASE, TOPV, PLOTH = 5.30, 9.0, 3.10
    def y_of(v):
        return BASE - v / TOPV * PLOTH

    STEPS = [('FY27', 0.0, 3.0, '+3%'), ('FY28', 3.0, 5.0, '+2%'),
             ('FY29', 5.0, 7.0, '+2%'), ('FY30', 7.0, 8.5, '+1.5%')]
    X0, BW, GAP = 1.05, 1.12, 0.42
    centers = []
    for i, (lab, lo, hi, delta) in enumerate(STEPS):
        cx = X0 + BW / 2 + i * (BW + GAP)
        centers.append(cx)
        yt, yb = y_of(hi), y_of(lo)
        shape(MSO_SHAPE.RECTANGLE, cx - BW / 2, yt, BW, yb - yt, fill=BRIGHT_BLUE)
        textbox(cx - BW / 2, yt - 0.30, BW, 0.26, [(delta, 12.5, True, BLUE)],
                align=PP_ALIGN.CENTER)
        if i:
            rule(centers[i - 1] - BW / 2, yb, (BW + GAP) + BW / 2, GREY_2,
                 0.75, dash=True)
        textbox(cx - BW / 2, BASE + 0.10, BW, 0.24, [(lab, 11, True, BLACK)],
                align=PP_ALIGN.CENTER)

    tcx = X0 + BW / 2 + 4 * (BW + GAP)
    shape(MSO_SHAPE.RECTANGLE, tcx - BW / 2, y_of(8.5), BW,
          BASE - y_of(8.5), fill=PURPLE)
    textbox(tcx - BW / 2, y_of(8.5) - 0.32, BW, 0.28,
            [('8 – 9%', 14, True, PURPLE)], align=PP_ALIGN.CENTER)
    textbox(tcx - BW / 2 - 0.20, BASE + 0.10, BW + 0.40, 0.24,
            [('FY30 total', 11, True, BLACK)], align=PP_ALIGN.CENTER)
    rule(centers[-1] - BW / 2, y_of(8.5), (BW + GAP) + BW / 2, GREY_2, 0.75,
         dash=True)
    rule(0.85, BASE, 7.90, GREY_3, 1.0)
    textbox(0.85, 1.72, 4.00, 0.24,
            [('Platform IP revenue, % of total', 10, False, GREY_1)])

    slab = shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.85, 5.66, 7.90, 0.44, adj=0.22)
    gradient(slab, PURPLE_DEEP, PURPLE)
    shape(MSO_SHAPE.OVAL, 1.06, 5.74, 0.28, 0.28, fill=WHITE)
    glyph('people', 1.20, 5.88, 0.16, PURPLE_DEEP)
    shape(MSO_SHAPE.OVAL, 8.26, 5.74, 0.28, 0.28, fill=WHITE)
    glyph('shield', 8.40, 5.88, 0.16, PURPLE_DEEP)
    textbox(0.85, 5.66, 7.90, 0.44,
            [('RESOURCE ARR FOUNDATION', 12, True, WHITE)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    PXr, PWr = 9.05, 3.63
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, PXr, 1.72, PWr, 4.38,
          fill=WHITE, line=GREY_3, adj=0.04)
    shape(MSO_SHAPE.OVAL, PXr + PWr / 2 - 0.34, 2.02, 0.68, 0.68, fill=BLUE)
    glyph('chart_up', PXr + PWr / 2, 2.36, 0.34, WHITE)
    textbox(PXr, 2.92, PWr, 0.60,
            [[('Platform IP revenue', 14, True, PURPLE)],
             [('layer grows', 14, True, PURPLE)]],
            align=PP_ALIGN.CENTER, line_spacing=1.0)
    rule(PXr + 0.60, 3.72, PWr - 1.20, GREY_3)
    textbox(PXr, 3.92, PWr, 0.62,
            [[('3%', 26, True, BLUE), ('   →   ', 16, True, GREY_2),
              ('8 – 9%', 26, True, PURPLE)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(PXr + 0.20, 4.62, PWr - 0.40, 0.24,
            [('FY27 → FY30', 10.5, False, GREY_1)], align=PP_ALIGN.CENTER)
    for i, hh in enumerate((0.30, 0.46, 0.66)):
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, PXr + 1.06 + i * 0.52, 5.62 - hh,
              0.36, hh, fill=LIGHT_BLUE if i < 2 else BLUE, adj=0.22)

    band = shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, 6.32, 12.03, 0.66, adj=0.14)
    gradient(band, PURPLE_DEEP, BLUE)
    shape(MSO_SHAPE.OVAL, 1.02, 6.47, 0.36, 0.36, fill=WHITE)
    glyph('bulb', 1.20, 6.65, 0.20, PURPLE)
    textbox(1.58, 6.32, 10.80, 0.66,
            [[('Resource ARR remains the foundation. ', 12.5, True, WHITE),
              ('Platform IP adds a growing transaction-based revenue stream.',
               12.5, False, PERIWINKLE)]],
            anchor=MSO_ANCHOR.MIDDLE)


# ═════════════════════════════════ 5 · Transforming Our People & Culture ═════
def slide_people():
    new_slide()
    header([('Transforming our ', BLACK), ('people & culture', PURPLE)],
           "Building DPO's talent engine for agentic AI.")

    CARDS = [('people', 'Management Talent Inclusion', '5+',
              'E7+ leaders hired for AI capability', PURPLE_DEEP),
             ('star', 'Elite Talent', '50+',
              'Offers to elite-vendor candidates, FY27', BLUE),
             ('chart_up', 'Resetting the Bar', '1,200+',
              'Assessed for technology aptitude', PURPLE),
             ('nodes', 'AI Capability', '53%+',
              'GenAI L1 completed', BLUE)]
    TW, TG, TT, TH = 2.82, 0.25, 1.70, 2.30
    for i, (icon, head, num, cap, col) in enumerate(CARDS):
        x = 0.65 + i * (TW + TG)
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, TT, TW, TH, fill=WHITE,
              line=GREY_3, adj=0.05)
        chip(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, x, TT, TW, 0.54,
             [[(head, 11.5, True, WHITE)]], fill=col, adj=(0.20, 0.0))
        shape(MSO_SHAPE.OVAL, x + TW / 2 - 0.28, TT + 0.68, 0.56, 0.56,
              fill=GREY_4)
        glyph(icon, x + TW / 2, TT + 0.96, 0.30, col)
        textbox(x, TT + 1.32, TW, 0.44, [(num, 25, True, col)],
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        rule(x + 0.62, TT + 1.82, TW - 1.24, GREY_3)
        textbox(x + 0.14, TT + 1.90, TW - 0.28, 0.34,
                [(cap, 9.5, False, BLACK)], align=PP_ALIGN.CENTER,
                line_spacing=1.05)

    BT, BH = 4.24, 1.66
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, BT, 12.03, BH, fill=GREY_4, adj=0.08)
    textbox(0.65, BT + 0.14, 12.03, 0.24,
            [('CULTURAL SHIFT REQUIRED FOR A STARTUP', 10.5, True, PURPLE)],
            align=PP_ALIGN.CENTER)
    VALUES = [('target', 'Own the Outcome'), ('people', 'Customer Obsessed'),
              ('bulb', 'Innovate Relentlessly'), ('union', 'One Team'),
              ('chart_up', 'Data Driven'), ('cubes', 'Built to Scale')]
    IX, IW = 1.00, 11.33
    step = IW / len(VALUES)
    for i, (icon, lab) in enumerate(VALUES):
        cx = IX + step / 2 + i * step
        shape(MSO_SHAPE.OVAL, cx - 0.27, BT + 0.52, 0.54, 0.54, fill=PURPLE)
        glyph(icon, cx, BT + 0.79, 0.28, WHITE)
        textbox(cx - step / 2 + 0.06, BT + 1.16, step - 0.12, 0.36,
                [(lab, 10, True, PURPLE_DEEP)], align=PP_ALIGN.CENTER,
                line_spacing=1.05)
        if i < len(VALUES) - 1:
            poly([(cx + step / 2 - 0.05, BT + 0.74),
                  (cx + step / 2 + 0.05, BT + 0.79),
                  (cx + step / 2 - 0.05, BT + 0.84)], BLUE, 1.2)

    band = shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, 6.14, 12.03, 0.82, adj=0.12)
    gradient(band, PURPLE_DEEP, BLUE)
    shape(MSO_SHAPE.OVAL, 1.10, 6.36, 0.44, 0.44, fill=WHITE)
    glyph('star', 1.32, 6.58, 0.26, PURPLE)
    vrule(1.78, 6.32, 0.46, PERIWINKLE)
    textbox(1.98, 6.14, 9.80, 0.82,
            [[('High Focus.  High Performance.  High Growth.', 15, True, WHITE)],
             [("Leadership immersion begins Sep'26", 11, False, PERIWINKLE)]],
            anchor=MSO_ANCHOR.MIDDLE, space_after=2)


slide_restructuring()
slide_fy30()
slide_growth()
slide_arr()
slide_people()
prs.save(OUT)
print('saved', OUT, '| slides:', len(prs.slides._sldIdLst))
