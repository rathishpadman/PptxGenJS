"""HCLTech 'Engineering in Action' case-study slides (BOI, Guardian Life)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TEMPLATE = '/root/.claude/skills/synced/hcltech-deck/assets/hcltech_base.pptx'
OUT = '/home/user/PptxGenJS/hcltech_engineering_in_action.pptx'

BLACK       = RGBColor(0x00, 0x00, 0x00)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
PURPLE      = RGBColor(0x5F, 0x1E, 0xBE)
PURPLE_DEEP = RGBColor(0x41, 0x14, 0x82)
LAVENDER    = RGBColor(0x8C, 0x69, 0xF0)
PERIWINKLE  = RGBColor(0xB9, 0xC8, 0xFF)
BLUE        = RGBColor(0x0F, 0x5F, 0xDC)
ICE_BLUE    = RGBColor(0xDC, 0xE6, 0xF0)
GREY_1      = RGBColor(0x82, 0x91, 0xA0)
GREY_3      = RGBColor(0xC8, 0xD2, 0xDD)
GREY_4      = RGBColor(0xE6, 0xEB, 0xF5)
FONT = 'Aptos'

prs = Presentation(TEMPLATE)
I = Inches
S = None          # current slide's shape tree, set per slide


def style(run, size, *, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def fill_tf(tf, paras, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            line_spacing=None, space_after=0):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, prun in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for text, size, bold, color in prun:
            style(p.add_run(), size, bold=bold, color=color)
            p.runs[-1].text = text


def textbox(x, y, w, h, runs, **kw):
    tb = S.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, runs if runs and isinstance(runs[0], list) else [runs], **kw)
    return tb


def shape(kind, x, y, w, h, *, fill=None, line=None, line_w=0.75, adj=None, rot=None):
    s = S.add_shape(kind, I(x), I(y), I(w), I(h))
    if adj is not None:
        vals = adj if isinstance(adj, (list, tuple)) else [adj]
        for i, v in enumerate(vals):
            if i < len(s.adjustments):
                s.adjustments[i] = v
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    if rot is not None:
        s.rotation = rot
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    return s


def poly(pts, color, width, close=False):
    b = S.build_freeform(I(pts[0][0]), I(pts[0][1]))
    b.add_line_segments([(I(x), I(y)) for x, y in pts[1:]], close=close)
    s = b.convert_to_shape()
    s.fill.background()
    s.line.color.rgb = color
    s.line.width = Pt(width)
    s.shadow.inherit = False
    return s


def rule(x, y, w, color, width=0.75):
    return poly([(x, y), (x + w, y)], color, width)


def label_box(x, y, w, h, text, size, color, *, bold=True, align=PP_ALIGN.CENTER,
              anchor=MSO_ANCHOR.TOP, line_spacing=1.05):
    return textbox(x, y, w, h, [(text, size, bold, color)],
                   align=align, anchor=anchor, line_spacing=line_spacing)


# ------------------------------------------------------------------ glyphs ---
def glyph(name, cx, cy, s, color):
    """Draw a vector glyph centred on (cx, cy) inside a box of side s."""
    u = s / 2.0                                   # half-extent
    lw = max(1.0, s * 2.2)                        # stroke weight in pt

    if name == 'target':
        shape(MSO_SHAPE.OVAL, cx - u, cy - u, s, s, line=color, line_w=lw)
        shape(MSO_SHAPE.OVAL, cx - u * .55, cy - u * .55, s * .55, s * .55,
              line=color, line_w=lw)
        shape(MSO_SHAPE.OVAL, cx - u * .18, cy - u * .18, s * .18, s * .18, fill=color)
    elif name == 'arrow_up':
        shape(MSO_SHAPE.UP_ARROW, cx - u * .8, cy - u * .8, s * .8, s * .8,
              fill=color, adj=(0.42, 0.42), rot=45)
    elif name == 'bulb':
        shape(MSO_SHAPE.OVAL, cx - u * .62, cy - u, s * .62, s * .62,
              line=color, line_w=lw)
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - u * .26, cy + u * .10,
              s * .26, s * .32, fill=color, adj=0.4)
    elif name in ('doc', 'checklist'):
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - u * .62, cy - u * .86,
              s * .62, s * .86, line=color, line_w=lw, adj=0.12)
        for i in range(3):
            y = cy - u * .48 + i * s * .21
            rule(cx - u * .34, y, s * .34, color, lw * .85)
        if name == 'checklist':
            poly([(cx + u * .16, cy + u * .30), (cx + u * .40, cy + u * .54),
                  (cx + u * .86, cy - u * .06)], color, lw * 1.1)
    elif name == 'database':
        for i in range(3):
            shape(MSO_SHAPE.OVAL, cx - u * .70, cy - u * .82 + i * s * .32,
                  s * .70, s * .30, line=color, line_w=lw)
    elif name == 'search':
        shape(MSO_SHAPE.OVAL, cx - u * .78, cy - u * .82, s * .60, s * .60,
              line=color, line_w=lw)
        poly([(cx + u * .14, cy + u * .16), (cx + u * .66, cy + u * .68)],
             color, lw * 1.2)
    elif name == 'person':
        shape(MSO_SHAPE.OVAL, cx - u * .28, cy - u * .82, s * .28, s * .28,
              line=color, line_w=lw)
        shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, cx - u * .58, cy - u * .16,
              s * .58, s * .46, line=color, line_w=lw, adj=(0.45, 0.0))
    elif name == 'people':
        for dx in (-u * .44, u * .10):
            shape(MSO_SHAPE.OVAL, dx + cx, cy - u * .74, s * .30, s * .30,
                  line=color, line_w=lw)
            shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, dx + cx - s * .05,
                  cy - u * .12, s * .40, s * .44, line=color, line_w=lw,
                  adj=(0.45, 0.0))
    elif name == 'shield':
        poly([(cx - u * .62, cy - u * .78), (cx + u * .62, cy - u * .78),
              (cx + u * .62, cy + u * .10), (cx, cy + u * .84),
              (cx - u * .62, cy + u * .10)], color, lw, close=True)
    elif name == 'frame':
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - u * .78, cy - u * .62,
              s * .78, s * .62, line=color, line_w=lw, adj=0.12)
        poly([(cx - u * .12, cy - u * .62), (cx - u * .12, cy + u * .00)],
             color, lw * .85)
    elif name == 'code':
        poly([(cx - u * .18, cy - u * .58), (cx - u * .72, cy),
              (cx - u * .18, cy + u * .58)], color, lw * 1.15)
        poly([(cx + u * .18, cy - u * .58), (cx + u * .72, cy),
              (cx + u * .18, cy + u * .58)], color, lw * 1.15)
    elif name == 'chart_up':
        for i, hh in enumerate((.34, .58, .84)):
            shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - u * .70 + i * s * .27,
                  cy + u * .62 - s * hh, s * .18, s * hh, fill=color, adj=0.3)
    elif name == 'cloud':
        shape(MSO_SHAPE.CLOUD, cx - u * .86, cy - u * .62, s * .86, s * .62,
              line=color, line_w=lw)
    elif name == 'nodes':
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - s * .10, cy - u * .74,
              s * .20, s * .20, fill=color, adj=0.25)
        for dx in (-.32, -.08, .16):
            shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + s * dx, cy + u * .28,
                  s * .16, s * .16, fill=color, adj=0.25)
        rule(cx - s * .24, cy + u * .08, s * .48, color, lw * .85)
    elif name == 'cubes':
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - s * .15, cy - u * .82,
              s * .30, s * .30, fill=color, adj=0.22)
        for dx in (-.38, .08):
            shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + s * dx, cy + u * .04,
                  s * .30, s * .30, fill=color, adj=0.22)
    else:
        raise ValueError('unknown glyph: %s' % name)


# ------------------------------------------------------------------ layout ---
LX, LW = 0.65, 6.55                    # left column
PX_, PW_, PY_, PH_ = 7.48, 5.20, 0.45, 6.50   # outcomes panel
CX_, CW_ = PX_ + 0.34, PW_ - 0.68              # panel content


def build(spec):
    global S
    slide = prs.slides.add_slide(prs.slide_layouts[16])   # Title Only
    S = slide.shapes

    # ---------------- left: eyebrow, title, subtitle ----------------
    eb = shape(MSO_SHAPE.ROUNDED_RECTANGLE, LX, 0.50, 2.05, 0.26,
               fill=GREY_4, adj=0.5)
    eb.text_frame.margin_left = eb.text_frame.margin_right = 0
    eb.text_frame.margin_top = eb.text_frame.margin_bottom = 0
    fill_tf(eb.text_frame, [[('ENGINEERING IN ACTION', 8.5, True, PURPLE)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    textbox(LX, 0.88, LW, 0.98,
            [[(t, 28, True, c) for t, c in spec['title']]], line_spacing=0.95)
    textbox(LX, 1.94, LW, 0.60, [(spec['subtitle'], 16, False, GREY_1)],
            line_spacing=1.1)

    # ---------------- left: challenge / transformation / solution ---
    CT, CH = 2.62, 1.98
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, LX, CT, LW, CH,
          fill=WHITE, line=GREY_3, adj=0.02)
    for i, (icon, label, text) in enumerate(spec['rows']):
        ry = CT + 0.10 + i * 0.60
        shape(MSO_SHAPE.OVAL, LX + 0.22, ry + 0.04, 0.42, 0.42, fill=GREY_4)
        glyph(icon, LX + 0.43, ry + 0.25, 0.24, PURPLE)
        textbox(LX + 0.78, ry + 0.09, 1.58, 0.32, [(label, 12, True, PURPLE)],
                anchor=MSO_ANCHOR.MIDDLE)
        poly([(LX + 2.44, ry + 0.06), (LX + 2.44, ry + 0.44)], GREY_3, 0.75)
        textbox(LX + 2.60, ry + 0.02, LW - 2.82, 0.48,
                [(text, 11, False, BLACK)], anchor=MSO_ANCHOR.MIDDLE,
                line_spacing=1.1)
        if i < 2:
            rule(LX + 0.22, CT + 0.68 + i * 0.60, LW - 0.44, GREY_4)

    # ---------------- left: numbered journey ------------------------
    label_box(LX, 4.82, 3.0, 0.24, spec['journey_label'], 11, PURPLE,
              align=PP_ALIGN.LEFT)
    step_w, gap = 1.10, 0.26
    for i, (icon, name) in enumerate(spec['steps']):
        cx = LX + step_w / 2 + i * (step_w + gap)
        badge = shape(MSO_SHAPE.OVAL, cx - 0.12, 5.16, 0.24, 0.24, fill=PURPLE)
        badge.text_frame.margin_left = badge.text_frame.margin_right = 0
        badge.text_frame.margin_top = badge.text_frame.margin_bottom = 0
        fill_tf(badge.text_frame, [[(str(i + 1), 9, True, WHITE)]],
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        shape(MSO_SHAPE.OVAL, cx - 0.31, 5.48, 0.62, 0.62, fill=GREY_4)
        glyph(icon, cx, 5.79, 0.32, PURPLE)
        label_box(cx - step_w / 2, 6.18, step_w, 0.44, name, 9.5, BLACK,
                  bold=False)
        if i < 4:
            ax = cx + step_w / 2 + gap / 2
            poly([(ax - 0.07, 5.79), (ax + 0.07, 5.79)], GREY_1, 1.0)
            poly([(ax + 0.02, 5.74), (ax + 0.07, 5.79), (ax + 0.02, 5.84)],
                 GREY_1, 1.0)

    # ---------------- right: outcomes panel -------------------------
    panel = shape(MSO_SHAPE.ROUNDED_RECTANGLE, PX_, PY_, PW_, PH_, adj=0.035)
    panel.fill.gradient()
    panel.fill.gradient_stops[0].color.rgb = PURPLE_DEEP
    panel.fill.gradient_stops[0].position = 0.0
    panel.fill.gradient_stops[1].color.rgb = PURPLE
    panel.fill.gradient_stops[1].position = 1.0
    panel.fill.gradient_angle = 45.0

    shape(MSO_SHAPE.OVAL, CX_, 0.72, 0.50, 0.50, fill=LAVENDER)
    glyph(spec['panel_icon'], CX_ + 0.25, 0.97, 0.26, WHITE)
    textbox(CX_ + 0.66, 0.72, CW_ - 0.66, 0.50,
            [('BUSINESS OUTCOMES', 18, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
    rule(CX_, 1.40, CW_, LAVENDER)

    head = spec['headline']
    if head['kind'] == 'caption':
        textbox(CX_, 1.54, CW_, 0.60, [(head['text'], 11, False, PERIWINKLE)],
                anchor=MSO_ANCHOR.MIDDLE)
    else:
        textbox(CX_, 1.52, 1.30, 0.66, [(head['big'], 36, True, WHITE)],
                anchor=MSO_ANCHOR.MIDDLE)
        poly([(CX_ + 1.34, 1.54), (CX_ + 1.34, 2.20)], LAVENDER, 0.75)
        textbox(CX_ + 1.50, 1.50, CW_ - 1.50, 0.70,
                [[(head['text'], 13, True, WHITE)],
                 [(head['sub'], 9.5, False, PERIWINKLE)]],
                anchor=MSO_ANCHOR.MIDDLE, space_after=2)

    tw, th = (CW_ - 0.12) / 2, 1.50
    for i, (num, lab, sub) in enumerate(spec['tiles']):
        tx = CX_ + (i % 2) * (tw + 0.12)
        ty = 2.38 + (i // 2) * (th + 0.15)
        textbox(tx, ty + 0.06, tw, 0.54, [(num, 34, True, WHITE)],
                anchor=MSO_ANCHOR.MIDDLE)
        textbox(tx, ty + 0.64, tw, 0.24, [(lab, 11.5, True, WHITE)])
        textbox(tx, ty + 0.90, tw - 0.10, 0.46, [(sub, 9, False, PERIWINKLE)],
                line_spacing=1.1)
    poly([(CX_ + tw + 0.06, 2.38), (CX_ + tw + 0.06, 5.53)], LAVENDER, 0.75)
    rule(CX_, 3.955, CW_, LAVENDER)

    rule(CX_, 5.70, CW_, LAVENDER)
    iw = CW_ / 4
    for i, (icon, name) in enumerate(spec['strip']):
        cx = CX_ + iw / 2 + i * iw
        glyph(icon, cx, 6.02, 0.34, WHITE)
        label_box(cx - iw / 2 + 0.04, 6.26, iw - 0.08, 0.44, name, 9.5, WHITE)
        if i < 3:
            poly([(CX_ + (i + 1) * iw, 5.86), (CX_ + (i + 1) * iw, 6.62)],
                 LAVENDER, 0.75)
    return slide


BOI = {
    'title': [('Chargeback cycle: ', BLACK), ('10 days to under 4', PURPLE)],
    'subtitle': 'BOI: Azure → AWS platform reset plus four orchestrated agents, '
                'running in production.',
    'rows': [
        ('target', 'Challenge',
         'Manual investigations across multiple systems'),
        ('arrow_up', 'Transformation',
         'Azure → AWS platform transition'),
        ('bulb', 'Solution',
         '4 agents orchestrated across GWS and Toscana'),
    ],
    'journey_label': 'HOW IT WORKS',
    'steps': [('doc', 'Intake'), ('database', 'AI Investigation'),
              ('search', 'Chargeback Analysis'), ('person', 'Review & Resolution'),
              ('shield', 'Audit & Reporting')],
    'panel_icon': 'chart_up',
    'headline': {'kind': 'caption',
                 'text': 'Illustrative impact — to be validated'},
    'tiles': [('62%', 'Faster Resolution', '~10 days → ~3.8 days'),
              ('55%', 'Analyst Productivity', 'Hours saved per analyst, weekly'),
              ('92%', 'First-Pass Accuracy', 'Less rework and follow-up'),
              ('$3.2M', 'Annualized Value', 'Cost avoidance from automation')],
    'strip': [('cloud', 'Scalable Platform'), ('nodes', 'Enterprise Integration'),
              ('shield', 'Built-in Governance'), ('arrow_up', 'Production Ready')],
}

GUARDIAN = {
    'title': [('First DPO agentic win — ', BLACK),
              ('and a reusable agent factory', PURPLE)],
    'subtitle': 'Guardian Life: from function-centric automation to an '
                'enterprise-wide agentic operating model.',
    'rows': [
        ('target', 'Challenge',
         'Fragmented operations across five business functions'),
        ('arrow_up', 'Transformation',
         'Function-centric automation → agentic operating model'),
        ('bulb', 'Solution',
         'DPO Platform — reusable agents, skills, governance'),
    ],
    'journey_label': 'OUR APPROACH',
    'steps': [('search', 'Discover'), ('checklist', 'Prioritize'),
              ('frame', 'Design'), ('code', 'Build'), ('chart_up', 'Scale')],
    'panel_icon': 'target',
    'headline': {'kind': 'hero', 'big': '#1',
                 'text': 'First DPO Platform Agentic Deal Win',
                 'sub': 'Proof of platform market readiness'},
    'tiles': [('39', 'Agent Recipes', 'Created and catalogued'),
              ('20+', 'Reusable Skills', 'Cross-domain skills library'),
              ('4', 'Business Domains', 'Benefits · Products · Wealth · Corporate'),
              ('3', 'Control Layers', 'Security · Risk · Compliance')],
    'strip': [('cubes', 'Platform Reusability'), ('people', 'Multi-Function Reach'),
              ('shield', 'Agent Governance'), ('chart_up', 'Repeatable Model')],
}

build(BOI)
build(GUARDIAN)
prs.save(OUT)
print('saved', OUT, '| slides:', len(prs.slides._sldIdLst))
