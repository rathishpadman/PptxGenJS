"""HCLTech slide: Engineering the Autonomy of Operations."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TEMPLATE = '/root/.claude/skills/synced/hcltech-deck/assets/hcltech_base.pptx'
OUT = '/home/user/PptxGenJS/hcltech_autonomy_of_operations.pptx'

BLACK       = RGBColor(0x00, 0x00, 0x00)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
PURPLE      = RGBColor(0x5F, 0x1E, 0xBE)
PURPLE_DEEP = RGBColor(0x41, 0x14, 0x82)
BLUE        = RGBColor(0x0F, 0x5F, 0xDC)
BRIGHT_BLUE = RGBColor(0x3C, 0x91, 0xFF)
ICE_BLUE    = RGBColor(0xDC, 0xE6, 0xF0)
GREY_1      = RGBColor(0x82, 0x91, 0xA0)
GREY_3      = RGBColor(0xC8, 0xD2, 0xDD)
GREY_4      = RGBColor(0xE6, 0xEB, 0xF5)
FONT = 'Aptos'

prs = Presentation(TEMPLATE)
slide = prs.slides.add_slide(prs.slide_layouts[16])
S = slide.shapes
I = Inches


def style(run, size, *, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def fill_tf(tf, paras, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            line_spacing=None, space_after=0):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, prun in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for text, size, bold, color in prun:
            style(p.add_run(), size, bold=bold, color=color)
            p.runs[-1].text = text


def textbox(x, y, w, h, runs, **kw):
    tb = S.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, runs if runs and isinstance(runs[0], list) else [runs], **kw)
    return tb


def shape(kind, x, y, w, h, *, fill=None, line=None, line_w=0.75, adj=None, rot=None):
    s = S.add_shape(kind, I(x), I(y), I(w), I(h))
    if adj is not None:
        vals = adj if isinstance(adj, (list, tuple)) else [adj]
        for i, v in enumerate(vals):
            if i < len(s.adjustments):
                s.adjustments[i] = v
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    if rot is not None:
        s.rotation = rot
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    return s


def gradient(s, c0, c1, angle=0.0):
    s.fill.gradient()
    s.fill.gradient_stops[0].color.rgb = c0
    s.fill.gradient_stops[0].position = 0.0
    s.fill.gradient_stops[1].color.rgb = c1
    s.fill.gradient_stops[1].position = 1.0
    s.fill.gradient_angle = angle
    return s


def poly(pts, color, width, close=False, fill=None):
    b = S.build_freeform(I(pts[0][0]), I(pts[0][1]))
    b.add_line_segments([(I(x), I(y)) for x, y in pts[1:]], close=close)
    s = b.convert_to_shape()
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    s.line.color.rgb = color
    s.line.width = Pt(width)
    s.shadow.inherit = False
    return s


def chip(kind, x, y, w, h, lines, *, fill, line=None, adj=0.5):
    c = shape(kind, x, y, w, h, fill=fill, line=line, adj=adj)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = I(0.04)
    tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, lines, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
            line_spacing=0.98)
    return c


# ------------------------------------------------------------------ glyphs ---
def g_people3(cx, cy, s, color):
    """Three-figure group, matching the reference's workforce mark."""
    lw = max(1.1, s * 3.0)
    shape(MSO_SHAPE.OVAL, cx - .115 * s, cy - .46 * s, .23 * s, .23 * s,
          line=color, line_w=lw)
    shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, cx - .21 * s, cy - .16 * s,
          .42 * s, .36 * s, line=color, line_w=lw, adj=(0.48, 0.0))
    for dx in (-.44, .25):
        shape(MSO_SHAPE.OVAL, cx + dx * s, cy - .38 * s, .19 * s, .19 * s,
              line=color, line_w=lw)
        shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, cx + (dx - .04) * s,
              cy - .10 * s, .27 * s, .30 * s, line=color, line_w=lw,
              adj=(0.48, 0.0))


def g_dollar_down(cx, cy, s, color):
    lw = max(1.1, s * 3.0)
    shape(MSO_SHAPE.OVAL, cx - .42 * s, cy - .46 * s, .84 * s, .84 * s,
          line=color, line_w=lw)
    textbox(cx - .42 * s, cy - .46 * s, .84 * s, .84 * s,
            [('$', s * 34, True, color)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    shape(MSO_SHAPE.DOWN_ARROW, cx + .26 * s, cy + .16 * s, .22 * s, .34 * s,
          fill=color, adj=(0.45, 0.45))


def g_chart_arrow(cx, cy, s, color):
    lw = max(1.1, s * 3.0)
    for i, hh in enumerate((.30, .46, .64)):
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - .46 * s + i * .30 * s,
              cy + .44 * s - hh * s, .18 * s, hh * s, fill=color, adj=0.28)
    poly([(cx - .30 * s, cy - .10 * s), (cx + .38 * s, cy - .48 * s)],
         color, lw)
    poly([(cx + .14 * s, cy - .46 * s), (cx + .40 * s, cy - .50 * s),
          (cx + .34 * s, cy - .24 * s)], color, lw, close=True, fill=color)


def g_target_cross(cx, cy, s, color):
    lw = max(1.1, s * 3.0)
    shape(MSO_SHAPE.OVAL, cx - .38 * s, cy - .38 * s, .76 * s, .76 * s,
          line=color, line_w=lw)
    shape(MSO_SHAPE.OVAL, cx - .17 * s, cy - .17 * s, .34 * s, .34 * s,
          line=color, line_w=lw)
    shape(MSO_SHAPE.OVAL, cx - .07 * s, cy - .07 * s, .14 * s, .14 * s, fill=color)
    for a, b in (((0, -.52), (0, -.34)), ((0, .34), (0, .52)),
                 ((-.52, 0), (-.34, 0)), ((.34, 0), (.52, 0))):
        poly([(cx + a[0] * s, cy + a[1] * s), (cx + b[0] * s, cy + b[1] * s)],
             color, lw)


def g_gear_cycle(cx, cy, s, color):
    lw = max(1.1, s * 2.6)
    shape(MSO_SHAPE.GEAR_6, cx - .30 * s, cy - .30 * s, .60 * s, .60 * s,
          line=color, line_w=lw)
    shape(MSO_SHAPE.OVAL, cx - .50 * s, cy - .50 * s, s, s, line=color, line_w=lw)
    # arrowheads tangent to the ring, reading as clockwise rotation
    r = .50 * s
    for ang_x, ang_y, rot in ((.707, -.707, 135), (-.707, .707, 315)):
        shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
              cx + ang_x * r - .13 * s, cy + ang_y * r - .12 * s,
              .26 * s, .24 * s, fill=color, rot=rot)


def divider(y, label, color):
    tb = textbox(0.65, y, 12.03, 0.24, [(label, 11, True, color)],
                 align=PP_ALIGN.CENTER)
    half = len(label) * 0.055
    poly([(1.20, y + 0.13), (6.665 - half - 0.28, y + 0.13)], GREY_3, 0.75)
    poly([(6.665 + half + 0.28, y + 0.13), (12.13, y + 0.13)], GREY_3, 0.75)
    return tb


# ------------------------------------------------------------------ header ---
textbox(0.65, 0.50, 12.03, 0.55,
        [[('Engineering the ', 28, True, BLACK),
          ('Autonomy of Operations', 28, True, PURPLE)]])
textbox(0.65, 1.12, 12.03, 0.32,
        [('Transforming labor-intensive operations into AI-native autonomous '
          'platforms.', 16, False, GREY_1)])

# -------------------------------------------------- transformation scale -----
divider(1.62, 'TRANSFORMATION SCALE', PURPLE)

CY, CH2 = 2.24, 1.86
for x, pill, col, num in ((1.34, 'TODAY', PURPLE, '~1,000'),
                          (6.34, 'FUTURE', BLUE, '~400')):
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, CY, 3.60, CH2,
          fill=WHITE, line=GREY_3, adj=0.05)
    g_people3(x + 1.80, 2.80, 0.62, col)
    textbox(x, 3.06, 3.60, 0.58, [(num, 36, True, col)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(x, 3.68, 3.60, 0.24, [('FTEs', 13, True, col)], align=PP_ALIGN.CENTER)
    chip(MSO_SHAPE.ROUNDED_RECTANGLE, x + 1.80 - 0.75, CY - 0.17, 1.50, 0.34,
         [[(pill, 10, True, WHITE)]], fill=col)

gradient(shape(MSO_SHAPE.RIGHT_ARROW, 5.22, 2.92, 0.86, 0.50,
               adj=(0.52, 0.42)), PURPLE_DEEP, PURPLE)
for i in range(3):
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, 4.99 + i * 0.075, 3.02, 0.04, 0.30,
          fill=GREY_3, adj=0.5)

shape(MSO_SHAPE.OVAL, 10.39, 2.37, 1.60, 1.60, fill=ICE_BLUE, line=BRIGHT_BLUE)
textbox(10.39, 2.72, 1.60, 0.44, [('~60%', 22, True, BLUE)],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
textbox(10.49, 3.16, 1.40, 0.44,
        [[('Workforce', 9.5, False, BLACK)], [('Optimization', 9.5, False, BLACK)]],
        align=PP_ALIGN.CENTER, line_spacing=1.05)

# ------------------------------------------------------- business outcomes ---
divider(4.34, 'BUSINESS OUTCOMES', PURPLE)

TW, TGAP, TY, TH = 2.82, 0.25, 4.74, 2.08
OUTCOMES = [
    (g_dollar_down, '30 – 40%', 'Operating Cost\nReduction', PURPLE_DEEP, GREY_4),
    (g_chart_arrow, '20 – 30%', 'Working Capital\nImprovement', PURPLE, GREY_4),
    (g_target_cross, '100 – 150 bps', 'Overall Margin\nImprovement', BLUE, ICE_BLUE),
    (g_gear_cycle, 'Long-term', 'Scalable & Autonomous\nOperations', BLUE, ICE_BLUE),
]
for i, (draw, num, label, col, tint) in enumerate(OUTCOMES):
    x = 0.65 + i * (TW + TGAP)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, TY, TW, TH,
          fill=WHITE, line=GREY_3, adj=0.05)
    shape(MSO_SHAPE.OVAL, x + TW / 2 - 0.39, TY + 0.22, 0.78, 0.78, fill=tint)
    draw(x + TW / 2, TY + 0.61, 0.62, col)
    textbox(x, TY + 1.14, TW, 0.46,
            [(num, 21 if len(num) < 12 else 18, True, col)],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    textbox(x + 0.12, TY + 1.62, TW - 0.24, 0.44,
            [[(ln, 10.5, False, BLACK)] for ln in label.split('\n')],
            align=PP_ALIGN.CENTER, line_spacing=1.05)

# --------------------------------------------------- brand bar (per the snap) -
gradient(shape(MSO_SHAPE.RECTANGLE, 0.0, 7.30, 13.333, 0.20), PURPLE_DEEP,
         BRIGHT_BLUE)

prs.save(OUT)
print('saved', OUT, '| shapes:', len(slide.shapes))
