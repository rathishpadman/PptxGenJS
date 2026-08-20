"""Build a single HCLTech-branded slide: portfolio diversification story."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

TEMPLATE = '/root/.claude/skills/synced/hcltech-deck/assets/hcltech_base.pptx'
OUT = '/home/user/PptxGenJS/hcltech_portfolio_diversification.pptx'

# --- locked HCL tokens -------------------------------------------------------
BLACK        = RGBColor(0x00, 0x00, 0x00)
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
PURPLE       = RGBColor(0x5F, 0x1E, 0xBE)
PURPLE_DEEP  = RGBColor(0x41, 0x14, 0x82)
LAVENDER     = RGBColor(0x8C, 0x69, 0xF0)
BLUE         = RGBColor(0x0F, 0x5F, 0xDC)
BRIGHT_BLUE  = RGBColor(0x3C, 0x91, 0xFF)
ICE_BLUE     = RGBColor(0xDC, 0xE6, 0xF0)
GREY_1       = RGBColor(0x82, 0x91, 0xA0)
GREY_2       = RGBColor(0xA5, 0xAF, 0xBE)
GREY_3       = RGBColor(0xC8, 0xD2, 0xDD)
GREY_4       = RGBColor(0xE6, 0xEB, 0xF5)
FONT = 'Aptos'

prs = Presentation(TEMPLATE)
slide = prs.slides.add_slide(prs.slide_layouts[17])   # Title w/t Subtitle
shapes = slide.shapes


def I(v):
    return Inches(v)


def style(run, size, *, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def textbox(x, y, w, h, runs, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            line_spacing=None, space_after=0):
    """runs: list of (text, size, bold, color) or list of such lists (paragraphs)."""
    tb = shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if runs and isinstance(runs[0], list) else [runs]
    for i, prun in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for text, size, bold, color in prun:
            style(p.add_run(), size, bold=bold, color=color)
            p.runs[-1].text = text
    return tb


def shape(kind, x, y, w, h, *, fill=None, line=None, line_w=0.75, adj=None, rot=None):
    s = shapes.add_shape(kind, I(x), I(y), I(w), I(h))
    if adj is not None:
        vals = adj if isinstance(adj, (list, tuple)) else [adj]
        for i, v in enumerate(vals):
            if i < len(s.adjustments):
                s.adjustments[i] = v
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    if rot is not None:
        s.rotation = rot
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    return s


def hline(x, y, w, color=GREY_3, width=0.75):
    from pptx.enum.shapes import MSO_CONNECTOR
    c = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x), I(y), I(x + w), I(y))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


# --- title + subtitle (brand placeholders) -----------------------------------
title = slide.shapes.title
title.text_frame.word_wrap = True
p = title.text_frame.paragraphs[0]
style(p.add_run(), 28, bold=True, color=BLACK)
p.runs[0].text = 'Diversifying the base is what expands the margin'

sub = [ph for ph in slide.placeholders if ph.placeholder_format.idx == 12][0]
sub.text_frame.word_wrap = True
sp = sub.text_frame.paragraphs[0]
style(sp.add_run(), 16, color=GREY_1)
sp.runs[0].text = ('25 → 80 accounts by FY30: lower concentration risk, '
                   '17%+ growth everywhere, 430 bps more EBIT')

# --- panel frame -------------------------------------------------------------
PY, PH, PW = 1.75, 4.20, 3.78
PX = [0.65, 4.775, 8.90]
HDR_H = 0.52


def panel(i, label, hdr_color):
    x = PX[i]
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, PY, PW, PH,
          fill=WHITE, line=GREY_3, adj=0.0265)
    shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, x, PY, PW, HDR_H,
          fill=hdr_color, adj=(0.192, 0.0))
    badge = shape(MSO_SHAPE.OVAL, x + 0.20, PY + 0.13, 0.26, 0.26, fill=WHITE)
    tf = badge.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    bp = tf.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    style(bp.add_run(), 11, bold=True, color=hdr_color)
    bp.runs[0].text = str(i + 1)
    textbox(x + 0.56, PY + 0.145, PW - 0.80, 0.24,
            [(label, 10, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
    return x


# ============================================================ panel 1: base ==
x = panel(0, 'BROADER BASE', PURPLE)
L = x + 0.28                      # left content edge
R = x + 1.92                      # right column edge

textbox(L, 2.42, 1.05, 0.72, [('25', 40, True, BLACK)], anchor=MSO_ANCHOR.MIDDLE)
textbox(x + 1.40, 2.42, 0.48, 0.72, [('→', 24, True, GREY_2)],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
textbox(R, 2.42, 1.30, 0.72, [('80', 40, True, PURPLE)], anchor=MSO_ANCHOR.MIDDLE)
textbox(L, 3.16, 1.10, 0.22, [('Accounts today', 9, False, GREY_1)])
textbox(R, 3.16, 1.55, 0.22, [('Accounts by FY30', 9, False, GREY_1)])

for dx, dy in [(0, 0), (0.34, 0), (0.68, 0), (0, 0.36), (0.34, 0.36)]:
    shape(MSO_SHAPE.OVAL, L + dx, 3.60 + dy, 0.26, 0.26, fill=PURPLE)
for row in range(4):
    for col in range(4):
        shape(MSO_SHAPE.OVAL, R + col * 0.21, 3.58 + row * 0.21, 0.14, 0.14,
              fill=LAVENDER)

hline(L, 4.75, 3.22)
textbox(L, 4.92, 3.22, 0.45,
        [('$24M', 20, True, BLACK), ('  →  ', 14, True, GREY_2),
         ('$15M', 20, True, PURPLE)], anchor=MSO_ANCHOR.MIDDLE)
textbox(L, 5.42, 3.22, 0.22, [('Average revenue per account', 9, False, GREY_1)])

# ================================================ panel 2: growth by vertical =
x = panel(1, 'WELL-DISTRIBUTED GROWTH', PURPLE)
L = x + 0.28
TRACK_X, TRACK_W = x + 1.64, 1.14
SCALE = 25.0

textbox(L, 2.38, 3.22, 0.22, [('FY26–FY30 revenue CAGR', 9, False, GREY_1)])
rows = [('Manufacturing', 23.4), ('Financial Services', 21.9),
        ('TMPE', 17.9), ('Mid Market', 17.2)]
for i, (name, val) in enumerate(rows):
    cy = 2.90 + i * 0.55
    textbox(L, cy - 0.13, 1.32, 0.26, [(name, 10, False, BLACK)],
            anchor=MSO_ANCHOR.MIDDLE)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, TRACK_X, cy - 0.10, TRACK_W, 0.20,
          fill=GREY_4, adj=0.5)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, TRACK_X, cy - 0.10,
          TRACK_W * val / SCALE, 0.20, fill=LAVENDER, adj=0.5)
    textbox(x + 2.84, cy - 0.13, 0.66, 0.26, [('%.1f%%' % val, 11, True, PURPLE)],
            align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

hline(L, 4.80, 3.22)
box = shape(MSO_SHAPE.ROUNDED_RECTANGLE, L, 4.97, 3.22, 0.75,
            fill=PURPLE, adj=0.09)
tf = box.text_frame
tf.margin_left = I(0.18); tf.margin_right = I(0.12)
tf.margin_top = tf.margin_bottom = 0
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p0 = tf.paragraphs[0]
p0.alignment = PP_ALIGN.LEFT
style(p0.add_run(), 22, bold=True, color=WHITE)
p0.runs[0].text = '19.5%'
p1 = tf.add_paragraph()
p1.alignment = PP_ALIGN.LEFT
style(p1.add_run(), 9, color=RGBColor(0xB9, 0xC8, 0xFF))
p1.runs[0].text = 'Portfolio CAGR, FY26–FY30'

# ====================================================== panel 3: EBIT returns =
x = panel(2, 'STRONGER RETURNS', BLUE)
L = x + 0.28

textbox(L, 2.45, 3.22, 0.92,
        [('+430', 54, True, BLUE), (' bps', 18, True, BLUE)],
        anchor=MSO_ANCHOR.MIDDLE)
textbox(L, 3.44, 3.22, 0.22,
        [('EBIT margin expansion, FY26 → FY30', 10, False, GREY_1)])

X0, X1 = x + 0.45, x + 3.35
YB, VMIN, VRANGE, PLOTH = 5.45, 12.5, 5.7, 1.30
pts = [(X0, 13.3), ((X0 + X1) / 2, 15.4), (X1, 17.6)]
xy = [(px, YB - (v - VMIN) / VRANGE * PLOTH) for px, v in pts]

area = shapes.build_freeform(I(xy[0][0]), I(YB))
area.add_line_segments([(I(px), I(py)) for px, py in xy] + [(I(X1), I(YB))],
                       close=True)
a = area.convert_to_shape()
a.fill.solid(); a.fill.fore_color.rgb = ICE_BLUE
a.line.fill.background(); a.shadow.inherit = False

ln = shapes.build_freeform(I(xy[0][0]), I(xy[0][1]))
ln.add_line_segments([(I(px), I(py)) for px, py in xy[1:]], close=False)
l = ln.convert_to_shape()
l.fill.background(); l.line.color.rgb = BLUE; l.line.width = Pt(2.25)
l.shadow.inherit = False

hline(X0 - 0.12, YB, (X1 - X0) + 0.24, color=GREY_3)
for px, py in xy:
    shape(MSO_SHAPE.OVAL, px - 0.065, py - 0.065, 0.13, 0.13,
          fill=BLUE, line=WHITE, line_w=1.25)

textbox(X0 - 0.02, xy[0][1] - 0.32, 0.70, 0.24, [('13.3%', 9, True, GREY_1)])
textbox(X1 - 0.70, xy[2][1] - 0.34, 0.70, 0.24, [('17.6%', 10, True, BLUE)],
        align=PP_ALIGN.RIGHT)
textbox(X0 - 0.02, YB + 0.08, 0.70, 0.22, [('FY26', 9, False, GREY_1)])
textbox(X1 - 0.70, YB + 0.08, 0.70, 0.22, [('FY30', 9, False, GREY_1)],
        align=PP_ALIGN.RIGHT)

# --- flow chevrons between panels --------------------------------------------
for cx in (4.60, 8.725):
    shape(MSO_SHAPE.ISOSCELES_TRIANGLE, cx - 0.10, 3.68, 0.20, 0.24,
          fill=GREY_2, rot=90)

# --- takeaway strip ----------------------------------------------------------
strip = shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, 6.18, 12.03, 0.58,
              fill=ICE_BLUE, adj=0.14)
tf = strip.text_frame
tf.margin_left = I(0.26); tf.margin_right = I(0.26)
tf.margin_top = tf.margin_bottom = 0
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
tp = tf.paragraphs[0]
tp.alignment = PP_ALIGN.LEFT
style(tp.add_run(), 12, bold=True, color=PURPLE_DEEP)
tp.runs[0].text = 'The trade is deliberate: '
style(tp.add_run(), 12, color=BLACK)
tp.runs[1].text = ('a smaller average account buys growth in every vertical '
                   'and 430 bps more EBIT.')

prs.save(OUT)
print('saved', OUT, '| shapes:', len(slide.shapes))
