"""Shared helpers and vector glyph library for the HCLTech board pack."""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BLACK       = RGBColor(0x00, 0x00, 0x00)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
PURPLE      = RGBColor(0x5F, 0x1E, 0xBE)
PURPLE_DEEP = RGBColor(0x41, 0x14, 0x82)
LAVENDER    = RGBColor(0x8C, 0x69, 0xF0)
PERIWINKLE  = RGBColor(0xB9, 0xC8, 0xFF)
BLUE        = RGBColor(0x0F, 0x5F, 0xDC)
BRIGHT_BLUE = RGBColor(0x3C, 0x91, 0xFF)
LIGHT_BLUE  = RGBColor(0x8C, 0xC8, 0xFA)
ICE_BLUE    = RGBColor(0xDC, 0xE6, 0xF0)
GREY_1      = RGBColor(0x82, 0x91, 0xA0)
GREY_2      = RGBColor(0xA5, 0xAF, 0xBE)
GREY_3      = RGBColor(0xC8, 0xD2, 0xDD)
GREY_4      = RGBColor(0xE6, 0xEB, 0xF5)
FONT = 'Aptos'
I = Inches

_S = None          # current slide shape tree


def use(shapes):
    global _S
    _S = shapes


def style(run, size, *, bold=False, color=BLACK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def fill_tf(tf, paras, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
            line_spacing=None, space_after=0):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, prun in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for text, size, bold, color in prun:
            style(p.add_run(), size, bold=bold, color=color)
            p.runs[-1].text = text


def textbox(x, y, w, h, runs, **kw):
    tb = _S.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, runs if runs and isinstance(runs[0], list) else [runs], **kw)
    return tb


def shape(kind, x, y, w, h, *, fill=None, line=None, line_w=0.75, adj=None,
          rot=None, dash=False):
    s = _S.add_shape(kind, I(x), I(y), I(w), I(h))
    if adj is not None:
        vals = adj if isinstance(adj, (list, tuple)) else [adj]
        for i, v in enumerate(vals):
            if i < len(s.adjustments):
                s.adjustments[i] = v
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if rot is not None:
        s.rotation = rot
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    return s


def hatch(x, y, w, h, fore, back=WHITE, line=None):
    s = _S.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), I(h))
    s.fill.patterned()
    s.fill.pattern = MSO_PATTERN.LIGHT_UPWARD_DIAGONAL
    s.fill.fore_color.rgb = fore
    s.fill.back_color.rgb = back
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    return s


def gradient(s, c0, c1, angle=0.0):
    s.fill.gradient()
    s.fill.gradient_stops[0].color.rgb = c0
    s.fill.gradient_stops[0].position = 0.0
    s.fill.gradient_stops[1].color.rgb = c1
    s.fill.gradient_stops[1].position = 1.0
    s.fill.gradient_angle = angle
    return s


def poly(pts, color, width, close=False, dash=False, fill=None):
    b = _S.build_freeform(I(pts[0][0]), I(pts[0][1]))
    b.add_line_segments([(I(x), I(y)) for x, y in pts[1:]], close=close)
    s = b.convert_to_shape()
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    s.line.color.rgb = color
    s.line.width = Pt(width)
    if dash:
        s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    s.shadow.inherit = False
    return s


def rule(x, y, w, color=GREY_3, width=0.75, dash=False):
    return poly([(x, y), (x + w, y)], color, width, dash=dash)


def vrule(x, y, h, color=GREY_3, width=0.75, dash=False):
    return poly([(x, y), (x, y + h)], color, width, dash=dash)


def arrowhead(cx, cy, color, size=0.07, direction='down'):
    d = {'down':  [(-size, -size * .9), (0, 0), (size, -size * .9)],
         'up':    [(-size, size * .9), (0, 0), (size, size * .9)],
         'right': [(-size * .9, -size), (0, 0), (-size * .9, size)]}[direction]
    return poly([(cx + a, cy + b) for a, b in d], color, 1.2, close=True, fill=color)


def chip(kind, x, y, w, h, lines, *, fill, line=None, adj=0.5, pad=0.06,
         align=PP_ALIGN.CENTER):
    c = shape(kind, x, y, w, h, fill=fill, line=line, adj=adj)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = I(pad)
    tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, lines, align=align, anchor=MSO_ANCHOR.MIDDLE, line_spacing=0.98)
    return c


# ------------------------------------------------------------------ glyphs ---
def glyph(name, cx, cy, s, color):
    u = s / 2.0
    lw = max(1.0, s * 2.4)
    if name == 'target':
        shape(MSO_SHAPE.OVAL, cx - u, cy - u, s, s, line=color, line_w=lw)
        shape(MSO_SHAPE.OVAL, cx - u * .5, cy - u * .5, s * .5, s * .5,
              line=color, line_w=lw)
        shape(MSO_SHAPE.OVAL, cx - u * .16, cy - u * .16, s * .16, s * .16, fill=color)
    elif name == 'people':
        shape(MSO_SHAPE.OVAL, cx - .115 * s, cy - .46 * s, .23 * s, .23 * s,
              line=color, line_w=lw)
        shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, cx - .21 * s, cy - .16 * s,
              .42 * s, .34 * s, line=color, line_w=lw, adj=(0.48, 0.0))
        for dx in (-.44, .25):
            shape(MSO_SHAPE.OVAL, cx + dx * s, cy - .38 * s, .19 * s, .19 * s,
                  line=color, line_w=lw)
            shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, cx + (dx - .04) * s,
                  cy - .10 * s, .27 * s, .28 * s, line=color, line_w=lw,
                  adj=(0.48, 0.0))
    elif name == 'chart_up':
        for i, hh in enumerate((.32, .52, .74)):
            shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - .48 * s + i * .32 * s,
                  cy + .40 * s - hh * s, .20 * s, hh * s, fill=color, adj=0.28)
    elif name == 'nodes':
        shape(MSO_SHAPE.OVAL, cx - .11 * s, cy - .46 * s, .22 * s, .22 * s, fill=color)
        for dx in (-.40, .18):
            shape(MSO_SHAPE.OVAL, cx + dx * s, cy + .20 * s, .22 * s, .22 * s,
                  fill=color)
        poly([(cx, cy - .20 * s), (cx - .29 * s, cy + .22 * s)], color, lw * .8)
        poly([(cx, cy - .20 * s), (cx + .29 * s, cy + .22 * s)], color, lw * .8)
        poly([(cx - .29 * s, cy + .30 * s), (cx + .29 * s, cy + .30 * s)],
             color, lw * .8)
    elif name == 'cubes':
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - .15 * s, cy - .48 * s,
              .30 * s, .30 * s, fill=color, adj=0.22)
        for dx in (-.38, .08):
            shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + dx * s, cy + .04 * s,
                  .30 * s, .30 * s, fill=color, adj=0.22)
    elif name == 'cube':
        shape(MSO_SHAPE.CUBE, cx - .42 * s, cy - .42 * s, .84 * s, .84 * s,
              line=color, line_w=lw)
    elif name == 'gear':
        shape(MSO_SHAPE.GEAR_6, cx - .46 * s, cy - .46 * s, .92 * s, .92 * s,
              line=color, line_w=lw)
        shape(MSO_SHAPE.OVAL, cx - .16 * s, cy - .16 * s, .32 * s, .32 * s,
              line=color, line_w=lw)
    elif name == 'shield':
        poly([(cx - .40 * s, cy - .48 * s), (cx + .40 * s, cy - .48 * s),
              (cx + .40 * s, cy + .06 * s), (cx, cy + .52 * s),
              (cx - .40 * s, cy + .06 * s)], color, lw, close=True)
    elif name == 'code':
        poly([(cx - .10 * s, cy - .38 * s), (cx - .46 * s, cy),
              (cx - .10 * s, cy + .38 * s)], color, lw)
        poly([(cx + .10 * s, cy - .38 * s), (cx + .46 * s, cy),
              (cx + .10 * s, cy + .38 * s)], color, lw)
    elif name == 'plane':
        poly([(cx - .44 * s, cy - .06 * s), (cx + .46 * s, cy - .40 * s),
              (cx + .06 * s, cy + .44 * s), (cx - .04 * s, cy + .08 * s)],
             color, lw, close=True, fill=color)
    elif name == 'layers':
        for i, dy in enumerate((-.30, -.02, .26)):
            shape(MSO_SHAPE.DIAMOND, cx - .44 * s, cy + dy * s - .13 * s,
                  .88 * s, .26 * s,
                  fill=color if i == 0 else None,
                  line=None if i == 0 else color, line_w=lw)
    elif name == 'star':
        shape(MSO_SHAPE.STAR_5_POINT, cx - .48 * s, cy - .48 * s,
              .96 * s, .96 * s, line=color, line_w=lw)
    elif name == 'bulb':
        shape(MSO_SHAPE.OVAL, cx - .30 * s, cy - .46 * s, .60 * s, .60 * s,
              line=color, line_w=lw)
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - .13 * s, cy + .16 * s,
              .26 * s, .30 * s, fill=color, adj=0.4)
    elif name == 'union':
        for dx in (-.18, .18):
            shape(MSO_SHAPE.OVAL, cx + dx * s - .30 * s, cy - .30 * s,
                  .60 * s, .60 * s, line=color, line_w=lw)
    elif name == 'trophy':
        shape(MSO_SHAPE.ROUND_2_DIAG_RECTANGLE, cx - .28 * s, cy - .44 * s,
              .56 * s, .48 * s, fill=color, adj=(0.5, 0.0))
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - .07 * s, cy + .04 * s,
              .14 * s, .22 * s, fill=color, adj=0.3)
        shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx - .30 * s, cy + .26 * s,
              .60 * s, .16 * s, fill=color, adj=0.4)
    elif name == 'dollar':
        shape(MSO_SHAPE.OVAL, cx - .44 * s, cy - .44 * s, .88 * s, .88 * s,
              line=color, line_w=lw)
        textbox(cx - .44 * s, cy - .44 * s, .88 * s, .88 * s,
                [('$', s * 32, True, color)],
                align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    else:
        raise ValueError('unknown glyph %r' % name)
