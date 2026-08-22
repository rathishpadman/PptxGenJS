"""HCLTech slide: Strengthening before we scale out."""
import sys
sys.path.insert(0, '.')
from pptx import Presentation
from hcltech_board_common import *          # noqa: F401,F403

TEMPLATE = '/root/.claude/skills/synced/hcltech-deck/assets/hcltech_base.pptx'
OUT = '/home/user/PptxGenJS/hcltech_strengthen_before_scale.pptx'

prs = Presentation(TEMPLATE)
slide = prs.slides.add_slide(prs.slide_layouts[16])
use(slide.shapes)


def lerp(a, b, t):
    return RGBColor(*(int(x + (y - x) * t) for x, y in zip(a, b)))


C0, C1 = (0x41, 0x14, 0x82), (0x0F, 0x5F, 0xDC)

textbox(0.65, 0.50, 12.03, 0.55,
        [[('Strengthening before ', 28, True, BLACK),
          ('we scale out', 28, True, PURPLE)]])
textbox(0.65, 1.12, 12.03, 0.32,
        [('Eight moves to define the organization — signed off before we launch.',
          16, False, GREY_1)])

# ---- SCALE OUT marker above the top rung ----------------------------------
chip(MSO_SHAPE.ROUNDED_RECTANGLE, 10.40, 1.40, 2.28, 0.50,
     [[('SCALE OUT', 14, True, BLUE)]], fill=ICE_BLUE, adj=0.22)
shape(MSO_SHAPE.UP_ARROW, 10.66, 1.53, 0.22, 0.24, fill=BLUE, adj=(0.42, 0.42))

# ---- the staircase: rung 1 at the foot, rung 8 at the top ------------------
RUNGS = [
    ('cubes',   'Strengthen capability units',
     'Build differentiated capabilities and functional depth'),
    ('shield',  'Transition and quality',
     'Execute transitions flawlessly and build a culture of quality'),
    ('layers',  'Delayering — reduced hierarchy',
     'Simplify structure to increase speed, agility and empowerment'),
    ('people',  'Span of control calibration',
     'Deliberate spans per layer so no manager is silently overloaded'),
    ('target',  'Decision rights and authority',
     'Who decides what, how fast, and where it escalates'),
    ('union',   'ISD in few accounts',
     'Drive integrated solution delivery in strategic accounts'),
    ('nodes',   'Transformation direct reporting',
     'Give transformation a clear mandate, visibility and accountability'),
    ('star',    'VDUs — customer alignment',
     'Align delivery units to customer outcomes and expectations'),
]
W, PITCH, RH, X0, STEP, TOP_Y = 9.40, 0.52, 0.48, 0.95, 0.30, 2.02
for i, (icon, title, desc) in enumerate(RUNGS):
    x = X0 + i * STEP
    y = TOP_Y + (len(RUNGS) - 1 - i) * PITCH
    col = lerp(C0, C1, i / (len(RUNGS) - 1))

    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, W, RH,
          fill=WHITE if i % 2 else GREY_4, line=GREY_3, adj=0.16)
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 0.07, RH, fill=col, adj=0.5)

    chip(MSO_SHAPE.OVAL, x + 0.18, y + 0.07, 0.34, 0.34,
         [[(str(i + 1), 12, True, WHITE)]], fill=col)
    textbox(x + 0.64, y, 2.80, RH, [(title, 10.5, True, col)],
            anchor=MSO_ANCHOR.MIDDLE)
    vrule(x + 3.50, y + 0.10, RH - 0.20, GREY_3)
    shape(MSO_SHAPE.OVAL, x + 3.66, y + 0.09, 0.30, 0.30, fill=WHITE, line=GREY_3)
    glyph(icon, x + 3.81, y + 0.24, 0.17, col)
    textbox(x + 4.10, y, W - 4.30, RH, [(desc, 9.5, False, BLACK)],
            anchor=MSO_ANCHOR.MIDDLE)

# ---- anchor in the wedge the staircase leaves open ------------------------
shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.68, 2.06, 0.07, 0.36, fill=PURPLE, adj=0.5)
textbox(0.90, 2.06, 2.00, 0.36,
        [[('STRONG TODAY.', 11, True, PURPLE_DEEP)],
         [('READY TOMORROW.', 11, True, PURPLE_DEEP)]], line_spacing=1.05)
textbox(0.90, 2.52, 1.68, 0.48,
        [('A solid foundation enables confident, sustainable scale.',
          9, False, GREY_1)], line_spacing=1.15)

# ---- closing band ----------------------------------------------------------
band = shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.65, 6.32, 12.03, 0.64, adj=0.14)
gradient(band, PURPLE_DEEP, BLUE)
shape(MSO_SHAPE.OVAL, 1.02, 6.47, 0.34, 0.34, fill=WHITE)
glyph('shield', 1.19, 6.64, 0.19, PURPLE_DEEP)
textbox(1.56, 6.32, 10.80, 0.64,
        [[('Strengthen the core. Define the organization. ', 13, True, WHITE),
          ('Then scale with confidence.', 13, True, PERIWINKLE)]],
        anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print('saved', OUT, '| shapes:', len(slide.shapes))
